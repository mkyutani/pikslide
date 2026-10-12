"""Check each fizzbuzz run's deck against the task, from the .pptx itself:
one slide; the twelve shapes, each of its kind and with its text; four
"yes" and four "no" labels; and nothing else on it, no other text, no
placeholder and no notes. What the deck can't tell -- that each arrow
goes where the task says, that lines don't cross and text doesn't
overlap -- is for looking at its PNG. Prints Markdown."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

RESULTS = Path(__file__).resolve().parent / "results" / "fizzbuzz"
METHODS = ("pikslide", "python-pptx", "raw-xml")
SHAPES = {
    "Start": "oval", "End": "oval", "i = 1": "rectangle", "i = i + 1": "rectangle",
    "i <= 100": "diamond", "i % 15 == 0": "diamond", "i % 3 == 0": "diamond", "i % 5 == 0": "diamond",
    "print i": "parallelogram", 'print "FizzBuzz"': "parallelogram", 'print "Fizz"': "parallelogram",
    'print "Buzz"': "parallelogram",
}
# Each kind's preset geometries: the plain one, or its flowchart symbol
KINDS = {"ellipse": "oval", "flowChartTerminator": "oval", "rect": "rectangle", "flowChartProcess": "rectangle",
         "diamond": "diamond", "flowChartDecision": "diamond",
         "parallelogram": "parallelogram", "flowChartInputOutput": "parallelogram"}
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def kind(shape) -> str | None:
    geometry = shape._element.find(f".//{_A}prstGeom")
    return None if geometry is None else KINDS.get(geometry.get("prst"), geometry.get("prst"))


def check(pptx: Path) -> list[str]:
    """What's wrong with the deck, if anything."""
    prs = Presentation(str(pptx))
    problems = [] if len(prs.slides) == 1 else [f"{len(prs.slides)} slides"]
    shapes = list(prs.slides[0].shapes)
    texts = {s.text_frame.text.strip(): kind(s) for s in shapes if s.has_text_frame and s.text_frame.text.strip()}
    problems += [f"{text!r}: {texts.get(text) or 'missing'}, not a {k}" for text, k in SHAPES.items()
                 if texts.get(text) != k]
    labels = [s.text_frame.text.strip() for s in shapes if s.has_text_frame and s.text_frame.text.strip() in ("yes", "no")]
    if (labels.count("yes"), labels.count("no")) != (4, 4):
        problems.append(f"{labels.count('yes')} yes and {labels.count('no')} no labels, not 4 and 4")
    problems += [f"other text {t!r}" for t in texts if t not in SHAPES and t not in ("yes", "no")]
    problems += [f"placeholder {s.name!r}" for s in shapes if s.is_placeholder]
    problems += ["notes" for slide in prs.slides if slide.has_notes_slide]
    return problems


def main() -> None:
    print("| method | run | deck |\n|---|---:|---|")
    for method in METHODS:
        for run in sorted((RESULTS / method).glob("[0-9][0-9]")):
            problems = check(run / "fizzbuzz.pptx")
            print(f"| {method} | {run.name} | {'; '.join(problems) or 'as the task says'} |")


if __name__ == "__main__":
    main()
