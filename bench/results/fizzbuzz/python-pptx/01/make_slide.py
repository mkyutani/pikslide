"""Write fizzbuzz.pptx: a one-slide flowchart of FizzBuzz for i = 1..100."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Arial"
INK = RGBColor(0x26, 0x32, 0x38)
LINE_W = Pt(1.5)

OVAL_FILL = RGBColor(0xD9, 0xEF, 0xD9)
RECT_FILL = RGBColor(0xDC, 0xE8, 0xF7)
DIAMOND_FILL = RGBColor(0xFF, 0xF2, 0xCC)
PARA_FILL = RGBColor(0xFB, 0xE3, 0xD4)
PARA_SLANT = 0.35  # parallelogram adjustment: slant = PARA_SLANT * height

# Geometry, in inches.
SLIDE_W, SLIDE_H = 13.333, 7.5
X_MAIN = 5.0  # centre of the main column
X_SIDE = 8.45  # centre of the prints and of End
X_LOOP = 3.1  # left-side return line
X_BUS = 10.2  # right-side line joining the prints
GAP = 0.32  # vertical space between shapes in the main column

SIZES = {
    "oval": (1.5, 0.48),
    "rect": (1.7, 0.48),
    "diamond": (2.6, 0.72),
    "para": (1.9, 0.48),
    "print": (2.3, 0.48),
}

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def emu(v):
    return Emu(int(round(Inches(v))))


def drop_style(shape):
    """Remove the theme style reference so only explicit formatting applies."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def style_line(line, arrow):
    line.color.rgb = INK
    line.width = LINE_W
    ln = line._get_or_add_ln()
    join = ln.makeelement(qn("a:miter"), {"lim": "800000"})
    ln.append(join)
    if arrow:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))


class Box:
    """A flowchart shape with its outline points, in inches."""

    def __init__(self, kind, text, cx, cy):
        self.kind, self.cx, self.cy = kind, cx, cy
        self.w, self.h = SIZES[kind]
        auto = {
            "oval": MSO_SHAPE.OVAL,
            "rect": MSO_SHAPE.RECTANGLE,
            "diamond": MSO_SHAPE.DIAMOND,
            "para": MSO_SHAPE.PARALLELOGRAM,
            "print": MSO_SHAPE.PARALLELOGRAM,
        }[kind]
        fill = {
            "oval": OVAL_FILL,
            "rect": RECT_FILL,
            "diamond": DIAMOND_FILL,
            "para": PARA_FILL,
            "print": PARA_FILL,
        }[kind]
        sp = slide.shapes.add_shape(auto, emu(cx - self.w / 2),
                                    emu(cy - self.h / 2), emu(self.w),
                                    emu(self.h))
        drop_style(sp)
        if kind in ("para", "print"):
            sp.adjustments[0] = PARA_SLANT
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
        style_line(sp.line, arrow=False)

        tf = sp.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(14)
        run.font.color.rgb = INK

    @property
    def top(self):
        return (self.cx, self.cy - self.h / 2)

    @property
    def bottom(self):
        return (self.cx, self.cy + self.h / 2)

    @property
    def left(self):
        # A parallelogram's left side is slanted; take its midpoint.
        inset = PARA_SLANT * self.h / 2 if self.kind in ("para", "print") else 0
        return (self.cx - self.w / 2 + inset, self.cy)

    @property
    def right(self):
        inset = PARA_SLANT * self.h / 2 if self.kind in ("para", "print") else 0
        return (self.cx + self.w / 2 - inset, self.cy)


def segment(p1, p2, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(p1[0]),
                                   emu(p1[1]), emu(p2[0]), emu(p2[1]))
    drop_style(c)
    style_line(c.line, arrow)
    return c


def polyline(points, arrow=True):
    """An open path through points (inches), with an arrowhead at the end."""
    scale = Inches(1) / 1000
    pts = [(round(x * 1000), round(y * 1000)) for x, y in points]
    builder = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=scale)
    builder.add_line_segments(pts[1:], close=False)
    shape = builder.convert_to_shape()
    drop_style(shape)
    shape.fill.background()
    style_line(shape.line, arrow)
    return shape


def label(text, cx, cy):
    w, h = 0.42, 0.22
    tb = slide.shapes.add_textbox(emu(cx - w / 2), emu(cy - h / 2), emu(w),
                                  emu(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = INK


# Main column, stacked top to bottom with an equal gap between shapes.
column = [
    ("oval", "Start"),
    ("rect", "i = 1"),
    ("diamond", "i <= 100"),
    ("diamond", "i % 15 == 0"),
    ("diamond", "i % 3 == 0"),
    ("diamond", "i % 5 == 0"),
    ("para", "print i"),
    ("rect", "i = i + 1"),
]
total_h = sum(SIZES[k][1] for k, _ in column) + GAP * (len(column) - 1)
y = (SLIDE_H - total_h) / 2
boxes = []
for kind, text in column:
    h = SIZES[kind][1]
    boxes.append(Box(kind, text, X_MAIN, y + h / 2))
    y += h + GAP
start, init, cond, d15, d3, d5, print_i, incr = boxes

# Down the column; label the arrows that leave a diamond.
down_labels = {cond: "yes", d15: "no", d3: "no", d5: "no"}
for upper, lower in zip(boxes, boxes[1:]):
    segment(upper.bottom, lower.top)
    if upper in down_labels:
        mid_y = (upper.bottom[1] + lower.top[1]) / 2
        label(down_labels[upper], X_MAIN + 0.27, mid_y)

# Loop finished: End, right of the loop test.
end = Box("oval", "End", X_SIDE, cond.cy)
segment(cond.right, end.left)
label("no", (cond.right[0] + end.left[0]) / 2, cond.cy - 0.16)

# Each divisibility test branches right to its own print.
prints = []
for diamond, text in ((d15, 'print "FizzBuzz"'), (d3, 'print "Fizz"'),
                      (d5, 'print "Buzz"')):
    pr = Box("print", text, X_SIDE, diamond.cy)
    segment(diamond.right, pr.left)
    label("yes", (diamond.right[0] + pr.left[0]) / 2, diamond.cy - 0.16)
    prints.append(pr)

# The prints join one line down the right side, into i = i + 1.
first, *rest = prints
polyline([first.right, (X_BUS, first.cy), (X_BUS, incr.cy), incr.right])
for pr in rest:
    segment(pr.right, (X_BUS, pr.cy), arrow=False)

# From i = i + 1 round the left side, back into the loop test.
polyline([incr.left, (X_LOOP, incr.cy), (X_LOOP, cond.cy), cond.left])

prs.save("fizzbuzz.pptx")
print("wrote fizzbuzz.pptx")
