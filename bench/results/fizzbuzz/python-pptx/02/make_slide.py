"""Write fizzbuzz.pptx: a single slide holding a FizzBuzz flowchart."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W, SLIDE_H = 13.333, 7.5  # inches

FONT = "Calibri"
TEXT_PT = 14
LABEL_PT = 12
INK = RGBColor(0x26, 0x32, 0x38)
LINE_PT = 1.5

FILL = {
    "oval": RGBColor(0xC8, 0xE6, 0xC9),
    "rect": RGBColor(0xBB, 0xDE, 0xFB),
    "diamond": RGBColor(0xFF, 0xF3, 0xC4),
    "para": RGBColor(0xE1, 0xBE, 0xE7),
}
GEOM = {
    "oval": MSO_SHAPE.OVAL,
    "rect": MSO_SHAPE.RECTANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "para": MSO_SHAPE.PARALLELOGRAM,
}

# Shape sizes (inches).
SIZE = {
    "oval": (1.4, 0.48),
    "rect": (1.5, 0.48),
    "diamond": (2.3, 0.66),
    "para": (1.6, 0.48),
}
PRINT_W = 2.3  # width of the three prints on the right

GAP_PLAIN = 0.27  # vertical gap between shapes with an unlabeled arrow
GAP_LABEL = 0.40  # vertical gap between shapes with a labeled arrow
YES_ARROW = 1.0  # length of the horizontal arrows out of the diamonds
TRUNK_GAP = 0.45  # prints' right edge to the joining line
LOOP_GAP = 0.5  # diamonds' left vertex to the loop-back line


def E(inches):
    return Emu(int(round(inches * 914400)))


class Box:
    """A placed shape: kind, center and size, in inches."""

    def __init__(self, kind, text, cx, cy, w=None):
        self.kind, self.text = kind, text
        dw, self.h = SIZE[kind]
        self.w = w if w is not None else dw
        self.cx, self.cy = cx, cy

    @property
    def top(self):
        return self.cy - self.h / 2

    @property
    def bottom(self):
        return self.cy + self.h / 2

    @property
    def left(self):
        """x of the outline's leftmost point at mid-height."""
        x = self.cx - self.w / 2
        if self.kind == "para":
            # The slanted edge is offset by 25% of the shorter side.
            x += min(self.w, self.h) * 0.25 / 2
        return x

    @property
    def right(self):
        """x of the outline's rightmost point at mid-height."""
        return 2 * self.cx - self.left


def style_text(tf, size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for run in p.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = INK


def draw_box(shapes, b):
    s = shapes.add_shape(
        GEOM[b.kind], E(b.cx - b.w / 2), E(b.top), E(b.w), E(b.h)
    )
    s.fill.solid()
    s.fill.fore_color.rgb = FILL[b.kind]
    s.line.color.rgb = INK
    s.line.width = Pt(LINE_PT)
    s.shadow.inherit = False
    s.text_frame.text = b.text
    style_text(s.text_frame, TEXT_PT)


def style_line(line, head):
    line.color.rgb = INK
    line.width = Pt(LINE_PT)
    if head:
        ln = line._get_or_add_ln()
        ln.append(
            ln.makeelement(
                qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}
            )
        )


def segment(shapes, p1, p2, head=True):
    c = shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, E(p1[0]), E(p1[1]), E(p2[0]), E(p2[1])
    )
    style_line(c.line, head)
    c.shadow.inherit = False


def polyline(shapes, points, head=True):
    fb = shapes.build_freeform(E(points[0][0]), E(points[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in points[1:]], close=False)
    s = fb.convert_to_shape()
    s.fill.background()
    style_line(s.line, head)
    s.shadow.inherit = False


def label(shapes, text, left, top, w, h, align, anchor):
    tb = shapes.add_textbox(E(left), E(top), E(w), E(h))
    tb.text_frame.text = text
    style_text(tb.text_frame, LABEL_PT, align, anchor)


LABEL_W, LABEL_H = 0.45, 0.24


def label_beside(shapes, text, x, y_top, y_bottom):
    """Label to the right of a vertical arrow at x."""
    y = (y_top + y_bottom) / 2
    label(shapes, text, x + 0.07, y - LABEL_H / 2, LABEL_W, LABEL_H,
          PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def label_above(shapes, text, x_left, x_right, y):
    """Label above a horizontal arrow at y."""
    x = (x_left + x_right) / 2
    label(shapes, text, x - LABEL_W / 2, y - 0.04 - LABEL_H, LABEL_W, LABEL_H,
          PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(SLIDE_W), E(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    shapes = slide.shapes

    # Horizontal extent: loop line on the left, joining line on the right.
    dia_half = SIZE["diamond"][0] / 2
    left_span = dia_half + LOOP_GAP
    right_span = dia_half + YES_ARROW + PRINT_W + TRUNK_GAP
    cx = SLIDE_W / 2 - (right_span - left_span) / 2

    # The main column, top to bottom, with the gap below each shape.
    column = [
        ("oval", "Start", GAP_PLAIN),
        ("rect", "i = 1", GAP_PLAIN),
        ("diamond", "i <= 100", GAP_LABEL),
        ("diamond", "i % 15 == 0", GAP_LABEL),
        ("diamond", "i % 3 == 0", GAP_LABEL),
        ("diamond", "i % 5 == 0", GAP_LABEL),
        ("para", "print i", GAP_PLAIN),
        ("rect", "i = i + 1", 0),
    ]
    total = sum(SIZE[k][1] + g for k, _, g in column)
    y = (SLIDE_H - total) / 2
    boxes = []
    for kind, text, gap in column:
        h = SIZE[kind][1]
        boxes.append(Box(kind, text, cx, y + h / 2))
        y += h + gap
    start, init, cond, d15, d3, d5, print_i, inc = boxes

    side_cx = cx + dia_half + YES_ARROW + PRINT_W / 2
    end = Box("oval", "End", 0, cond.cy)
    end.cx = cx + dia_half + YES_ARROW + end.w / 2
    prints = [
        Box("para", 'print "FizzBuzz"', side_cx, d15.cy, PRINT_W),
        Box("para", 'print "Fizz"', side_cx, d3.cy, PRINT_W),
        Box("para", 'print "Buzz"', side_cx, d5.cy, PRINT_W),
    ]

    for b in boxes + [end] + prints:
        draw_box(shapes, b)

    # Down the column.
    for upper, lower in zip(boxes, boxes[1:]):
        segment(shapes, (cx, upper.bottom), (cx, lower.top))
    label_beside(shapes, "yes", cx, cond.bottom, d15.top)
    for upper, lower in [(d15, d3), (d3, d5), (d5, print_i)]:
        label_beside(shapes, "no", cx, upper.bottom, lower.top)

    # i <= 100 --no--> End
    segment(shapes, (cond.right, cond.cy), (end.left, end.cy))
    label_above(shapes, "no", cond.right, end.left, cond.cy)

    # Each divisibility test --yes--> its print.
    for d, p in zip([d15, d3, d5], prints):
        segment(shapes, (d.right, d.cy), (p.left, p.cy))
        label_above(shapes, "yes", d.right, p.left, d.cy)

    # The prints join one line down the right side, into i = i + 1.
    trunk_x = prints[0].right + TRUNK_GAP
    polyline(shapes, [
        (prints[0].right, prints[0].cy),
        (trunk_x, prints[0].cy),
        (trunk_x, inc.cy),
        (inc.right, inc.cy),
    ])
    for p in prints[1:]:
        segment(shapes, (p.right, p.cy), (trunk_x, p.cy), head=False)

    # i = i + 1 loops round the left side, back into i <= 100.
    loop_x = cond.cx - dia_half - LOOP_GAP
    polyline(shapes, [
        (inc.left, inc.cy),
        (loop_x, inc.cy),
        (loop_x, cond.cy),
        (cond.left, cond.cy),
    ])

    prs.save("fizzbuzz.pptx")


if __name__ == "__main__":
    main()
