"""Write fizzbuzz.pptx: a single slide holding a FizzBuzz flowchart."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Arial"
INK = RGBColor(0x26, 0x32, 0x38)
LINE_WIDTH = Pt(1.5)

FILLS = {
    "oval": RGBColor(0xC8, 0xE6, 0xC9),
    "rect": RGBColor(0xBB, 0xDE, 0xFB),
    "diamond": RGBColor(0xFF, 0xF1, 0xB8),
    "para": RGBColor(0xFF, 0xD9, 0xB3),
}
KINDS = {
    "oval": MSO_SHAPE.OVAL,
    "rect": MSO_SHAPE.RECTANGLE,
    "diamond": MSO_SHAPE.DIAMOND,
    "para": MSO_SHAPE.PARALLELOGRAM,
}

# Geometry, in inches.
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_Y = 0.25
BOX_H = 0.45  # ovals, rectangles and parallelograms
DIAMOND_W, DIAMOND_H = 2.4, 0.72
RECT_W = 1.5
OVAL_W = 1.2
PARA_W = 2.2
PARA_ADJ = 0.5  # slant = PARA_ADJ * BOX_H
YES_RUN = 1.1  # gap between a diamond's right vertex and its print
LOOP_GAP = 0.45  # gap between the widest column shape and the loop line
BUS_GAP = 0.45  # gap between the prints and the line down the right side

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
shapes = slide.shapes


def emu(v):
    return Emu(int(round(Inches(v))))


def style_text(text_frame, text, size):
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    text_frame.word_wrap = False
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = INK


class Node:
    """A flowchart shape, centred on (cx, cy)."""

    def __init__(self, kind, text, cx, cy, w, h):
        self.kind, self.cx, self.cy, self.w, self.h = kind, cx, cy, w, h
        shape = shapes.add_shape(
            KINDS[kind], emu(cx - w / 2), emu(cy - h / 2), emu(w), emu(h)
        )
        if kind == "para":
            shape.adjustments[0] = PARA_ADJ
        shape.fill.solid()
        shape.fill.fore_color.rgb = FILLS[kind]
        shape.line.color.rgb = INK
        shape.line.width = LINE_WIDTH
        shape.shadow.inherit = False
        style_text(shape.text_frame, text, 14)

    @property
    def top(self):
        return (self.cx, self.cy - self.h / 2)

    @property
    def bottom(self):
        return (self.cx, self.cy + self.h / 2)

    @property
    def left(self):
        # A parallelogram's left side is slanted; its midpoint sits half a
        # slant in from the bounding box.
        inset = PARA_ADJ * self.h / 2 if self.kind == "para" else 0
        return (self.cx - self.w / 2 + inset, self.cy)

    @property
    def right(self):
        inset = PARA_ADJ * self.h / 2 if self.kind == "para" else 0
        return (self.cx + self.w / 2 - inset, self.cy)


def style_line(shape, arrow):
    shape.line.color.rgb = INK
    shape.line.width = LINE_WIDTH
    ln = shape.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:round"))
    if arrow:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


def path(points, arrow=True):
    """A line through points (inches), with an arrowhead at the last one."""
    if len(points) == 2:
        (x0, y0), (x1, y1) = points
        shape = shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, emu(x0), emu(y0), emu(x1), emu(y1)
        )
    else:
        x0, y0 = points[0]
        builder = shapes.build_freeform(emu(x0), emu(y0), scale=1.0)
        builder.add_line_segments([(emu(x), emu(y)) for x, y in points[1:]],
                                  close=False)
        shape = builder.convert_to_shape()
        shape.fill.background()
    shape.shadow.inherit = False
    style_line(shape, arrow)
    return shape


def label(text, x, y, anchor):
    """A 'yes'/'no' label; anchor says which corner of it sits at (x, y)."""
    w, h = 0.36, 0.22
    left = x if "left" in anchor else x - w
    top = y if "top" in anchor else y - h
    box = shapes.add_textbox(emu(left), emu(top), emu(w), emu(h))
    style_text(box.text_frame, text, 12)
    box.text_frame.paragraphs[0].alignment = (
        PP_ALIGN.LEFT if "left" in anchor else PP_ALIGN.RIGHT
    )


# The main column, top to bottom.
column = [
    ("oval", "Start", OVAL_W, BOX_H),
    ("rect", "i = 1", RECT_W, BOX_H),
    ("diamond", "i <= 100", DIAMOND_W, DIAMOND_H),
    ("diamond", "i % 15 == 0", DIAMOND_W, DIAMOND_H),
    ("diamond", "i % 3 == 0", DIAMOND_W, DIAMOND_H),
    ("diamond", "i % 5 == 0", DIAMOND_W, DIAMOND_H),
    ("para", "print i", PARA_W, BOX_H),
    ("rect", "i = i + 1", RECT_W, BOX_H),
]
gap = (SLIDE_H - 2 * MARGIN_Y - sum(h for *_, h in column)) / (len(column) - 1)

# Centre the whole drawing, from the loop line to the right-hand line.
left_extent = DIAMOND_W / 2 + LOOP_GAP
right_extent = DIAMOND_W / 2 + YES_RUN + PARA_W + BUS_GAP
cx = SLIDE_W / 2 - (right_extent - left_extent) / 2
side_cx = cx + DIAMOND_W / 2 + YES_RUN + PARA_W / 2  # prints and End

nodes = []
y = MARGIN_Y
for kind, text, w, h in column:
    nodes.append(Node(kind, text, cx, y + h / 2, w, h))
    y += h + gap
start, init, cond, mod15, mod3, mod5, print_i, incr = nodes

# Down the column.
for a, b in zip(nodes, nodes[1:]):
    path([a.bottom, b.top])
LX, LY = 0.08, 0.04  # label offset from its line
label("yes", cond.bottom[0] + LX, cond.bottom[1] + LY, "top-left")
for d in (mod15, mod3, mod5):
    label("no", d.bottom[0] + LX, d.bottom[1] + LY, "top-left")

# i <= 100 fails: End.
end = Node("oval", "End", side_cx, cond.cy, OVAL_W, BOX_H)
path([cond.right, end.left])
label("no", cond.right[0] + 0.12, cond.cy - LY, "bottom-left")

# Each divisibility test that passes goes right to its own print.
prints = []
for d, text in ((mod15, 'print "FizzBuzz"'), (mod3, 'print "Fizz"'),
                (mod5, 'print "Buzz"')):
    p = Node("para", text, side_cx, d.cy, PARA_W, BOX_H)
    path([d.right, p.left])
    label("yes", d.right[0] + 0.12, d.cy - LY, "bottom-left")
    prints.append(p)

# The prints merge into one line down the right side, into i = i + 1.
bus_x = side_cx + PARA_W / 2 + BUS_GAP
path([prints[0].right, (bus_x, prints[0].cy), (bus_x, incr.cy), incr.right])
for p in prints[1:]:
    path([p.right, (bus_x, p.cy)], arrow=False)

# i = i + 1 loops round the left side, back into i <= 100.
loop_x = cx - left_extent
path([incr.left, (loop_x, incr.cy), (loop_x, cond.cy), cond.left])

prs.save("fizzbuzz.pptx")
print("wrote fizzbuzz.pptx")
