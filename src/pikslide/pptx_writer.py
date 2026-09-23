"""Render a resolved pik layout (:mod:`pikslide.pik.layout`) to a PowerPoint file.

Coordinates in a :class:`~pikslide.pik.layout.LayoutResult` are inches with
y pointing up (pikchr's convention); PowerPoint slides use EMU with y
pointing down from the top-left, so this module flips y and adds a margin
around the diagram's bounding box.

Shape-kind to PowerPoint mapping is necessarily approximate for a few
classes (cylinder, file) since python-pptx's autoshape set doesn't have
an exact equivalent; see the per-kind comments below.
"""

from __future__ import annotations

import io
import os
import tempfile
import zipfile

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.image import ImagePart
from pptx.util import Emu, Inches, Pt

from .pik import ast
from .pik.layout import (
    NOT_RENDERED,
    Color,
    LayoutError,
    LayoutResult,
    Shape,
    _did_you_mean,
    _text_size_name,
    assign_text_slots,
    default_text_sizes,
    flatten_shapes,
    rasterize_svg,
    resolve_layout,
)

EMU_PER_INCH = 914400

# Vertical offset (in line-steps, positive = up in pik's y-up space) for
# each text slot pik_txt_vertical_layout() can assign -- see assign_text_slots().
_SLOT_STEP = {"above2": 2, "above": 1, "center": 0, "below": -1, "below2": -2}

# What "fit" sizing actually measures text *with* (a real installed font
# file, via Pillow) is necessarily independent of the font family the
# .pptx itself requests: by default that's a symbolic theme reference,
# not a literal name at all (docs/spec.md SS3.3, see _apply_run_font()
# below), so there is no one "the requested font" to look up a substitute
# for even in principle. PilFontMetrics instead tries a few widely-
# available, metrically-reasonable substitutes, in order.
#
# A Latin-only substitute (Liberation/DejaVu/Arial) silently *undersizes*
# any CJK text: missing-glyph fallback advances are far narrower than a
# real ideograph, so "fit" shapes come out too small and the text overflows
# them once PowerPoint actually renders it with a CJK-capable font. A Noto
# Sans CJK file, where present, covers Latin *and* CJK correctly, so it's
# tried first; the Latin-only substitutes remain as a fallback chain for
# machines without it (where only non-CJK text will still measure well).
_MEASURE_FONT_CANDIDATES = [
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    "/System/Library/Fonts/Supplemental/NotoSansCJKjp-Regular.otf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "C:\\Windows\\Fonts\\arial.ttf",
]

_LABEL_STEP_IN = 0.10  # how far a label shifts per above/below slot step at the medium size -- see _SLOT_STEP


def _find_measure_font() -> str | None:
    for path in _MEASURE_FONT_CANDIDATES:
        if os.path.exists(path):
            return path
    return None


class PilFontMetrics:
    """A pikslide.pik.layout.FontMetrics backed by real glyph widths (via
    Pillow), so a "fit" object's size tracks its actual text content --
    unlike a fixed per-character-width guess, this gets more (not less)
    accurate as text gets longer or more varied.

    Measures with whatever font this machine actually has installed among
    _MEASURE_FONT_CANDIDATES -- an approximation independent of what the
    .pptx itself actually requests, which by default is a symbolic theme
    font reference, not a literal name at all (docs/spec.md SS3.3, see
    _apply_run_font() below), so on a viewer whose theme font isn't
    metrically close to the substitute, "fit" sizing can disagree with
    what's actually rendered. This still generalizes correctly across
    different text content, unlike a flat estimate.
    """

    def __init__(
        self,
        text_sizes: dict[str, float] | None = None,
        font_path: str | None = None,
    ):
        # inches, keyed "small"/"medium"/"large" -- docs/spec.md SS3.3/SS3.7.
        self._text_sizes = text_sizes if text_sizes is not None else default_text_sizes()
        self._font_path = font_path if font_path is not None else _find_measure_font()
        self._cache: dict[int, ImageFont.FreeTypeFont] = {}

    def _size_pt(self, flags: list[str], text_sizes: dict[str, float] | None) -> float:
        # An explicit override (a specific object's own Shape.text_sizes,
        # docs/spec.md SS3.3, ext) takes precedence over this instance's
        # own default, so measuring an already-placed "fit" object can't
        # be skewed by a `medium`/etc. override later in the document.
        sizes = text_sizes if text_sizes is not None else self._text_sizes
        medium = sizes.get("medium", 10.5 / 72.0)
        inches = sizes.get(_text_size_name(flags), medium)
        return inches * 72.0

    def _font(
        self, flags: list[str], text_sizes: dict[str, float] | None
    ) -> tuple[ImageFont.FreeTypeFont | None, float]:
        size_pt = self._size_pt(flags, text_sizes)
        size_px = max(1, round(size_pt))
        if size_px not in self._cache:
            self._cache[size_px] = ImageFont.truetype(self._font_path, size_px) if self._font_path else None
        return self._cache[size_px], size_pt

    def text_width(self, text: str, flags: list[str] = (), text_sizes: dict[str, float] | None = None) -> float:
        font, size_pt = self._font(flags, text_sizes)
        if font is not None:
            px = font.getlength(text) if text else 0.0
        else:
            # No real font file found on this machine at all: fall back to
            # a flat estimate rather than failing outright.
            px = len(text) * size_pt * 0.55
        return px / 72.0 + (size_pt * 0.5) / 72.0

    def line_height(self, flags: list[str] = (), text_sizes: dict[str, float] | None = None) -> float:
        _font, size_pt = self._font(flags, text_sizes)
        return size_pt / 72.0


def resolve_for_pptx(
    doc: ast.Document,
    base_dir: str = ".",
    settings_text: str | None = None,
    settings_base_dir: str = ".",
) -> LayoutResult:
    """resolve_layout(), using real font metrics so "fit" objects are
    sized to match what write_pptx() will actually draw. Text sizes come
    from the prelude's own small/medium/large (docs/spec.md SS3.7), not a
    caller-supplied constant; write_pptx() then reads the *same* sizes
    back from the returned LayoutResult, so the two always agree.

    `base_dir` is the source file's own directory, against which an
    `image` object's path resolves (docs/spec.md SS3.5). `settings_text`/
    `settings_base_dir` are a template's settings file (docs/spec.md
    SS3.8), if any -- see find_settings_file()."""
    return resolve_layout(
        doc, metrics=PilFontMetrics(), base_dir=base_dir, settings_text=settings_text, settings_base_dir=settings_base_dir
    )


def _font_size(flags: list[str], text_sizes: dict[str, float]) -> Pt:
    medium = text_sizes.get("medium", 10.5 / 72.0)
    inches = text_sizes.get(_text_size_name(flags), medium)
    return Pt(inches * 72.0)


_AUTOSHAPE = {
    "box": MSO_SHAPE.RECTANGLE,
    "circle": MSO_SHAPE.OVAL,
    "ellipse": MSO_SHAPE.OVAL,
    "oval": MSO_SHAPE.OVAL,
    "diamond": MSO_SHAPE.DIAMOND,
    "cylinder": MSO_SHAPE.CAN,  # closest built-in equivalent
    "file": MSO_SHAPE.FLOWCHART_DOCUMENT,  # closest built-in equivalent
    "dot": MSO_SHAPE.OVAL,
    "text": MSO_SHAPE.RECTANGLE,  # rendered with no fill/outline, see below
}


def _rgb(value: int) -> RGBColor:
    v = value & 0xFFFFFF
    return RGBColor((v >> 16) & 0xFF, (v >> 8) & 0xFF, v & 0xFF)


# docs/spec.md SS3.3: a Color's theme_slot is one of these OOXML schemeClr
# reference names (see pik.layout.THEME_SLOTS for the full explanation of
# why these, and not dk1/lt1/dk2/lt2, are the ones a shape can reference).
_MSO_THEME_COLOR = {
    "tx1": MSO_THEME_COLOR.TEXT_1, "bg1": MSO_THEME_COLOR.BACKGROUND_1,
    "tx2": MSO_THEME_COLOR.TEXT_2, "bg2": MSO_THEME_COLOR.BACKGROUND_2,
    "accent1": MSO_THEME_COLOR.ACCENT_1, "accent2": MSO_THEME_COLOR.ACCENT_2,
    "accent3": MSO_THEME_COLOR.ACCENT_3, "accent4": MSO_THEME_COLOR.ACCENT_4,
    "accent5": MSO_THEME_COLOR.ACCENT_5, "accent6": MSO_THEME_COLOR.ACCENT_6,
    "hlink": MSO_THEME_COLOR.HYPERLINK, "folHlink": MSO_THEME_COLOR.FOLLOWED_HYPERLINK,
}


def _brightness(color: Color) -> float:
    """The python-pptx `ColorFormat.brightness` (-1..1) that reproduces a
    Color's lum_mod/lum_off (checked: brightness > 0 writes lumMod=
    (1-b)*100000/lumOff=b*100000, matching "Lighter N%"; brightness < 0
    writes lumMod=(1+b)*100000 alone, matching "Darker N%")."""
    if color.lum_off > 0:
        return color.lum_off
    if color.lum_mod < 1.0:
        return color.lum_mod - 1.0
    return 0.0


def _apply_color(color_format, color: Color) -> None:
    """Set a python-pptx ColorFormat (`fill.fore_color`, `line.color`, or
    `font.color` -- the same API for all three, checked) from a Color: a
    theme color is written as `schemeClr` (never resolved to RGB, so it
    stays linked to whatever theme the diagram lands in), an RGB color as
    `srgbClr`; `lighter`/`darker` apply to either the same way."""
    if color.theme_slot is not None:
        color_format.theme_color = _MSO_THEME_COLOR[color.theme_slot]
    else:
        color_format.rgb = _rgb(color.rgb or 0)
    b = _brightness(color)
    if b != 0.0:
        color_format.brightness = b


_MIN_SLIDE_SIDE_IN = 1.0  # PowerPoint refuses a slide smaller than 1in on either side


class _Transform:
    """Maps pik inches (y-up) to slide inches (y-down), with a margin.

    PowerPoint requires each slide dimension to be at least 1in; a small
    diagram (or a very small margin) can fall under that on its own, so
    slide_width/slide_height are clamped up to it, and the extra space is
    split evenly as additional margin so the diagram stays centered
    rather than pinned in a corner.
    """

    def __init__(self, bbox: tuple[float, float, float, float], margin: float):
        self.x0, self.y0, self.x1, self.y1 = bbox
        natural_w = (self.x1 - self.x0) + 2 * margin
        natural_h = (self.y1 - self.y0) + 2 * margin
        self.slide_width = max(natural_w, _MIN_SLIDE_SIDE_IN)
        self.slide_height = max(natural_h, _MIN_SLIDE_SIDE_IN)
        self.margin_x = margin + (self.slide_width - natural_w) / 2
        self.margin_y = margin + (self.slide_height - natural_h) / 2

    def rect(self, shape: Shape) -> tuple[float, float, float, float]:
        """Return (left, top, width, height) in inches for shape's bbox."""
        bx0, by0, bx1, by1 = shape.bbox
        left = (bx0 - self.x0) + self.margin_x
        top = (self.y1 - by1) + self.margin_y
        return left, top, bx1 - bx0, by1 - by0

    def point(self, pt: tuple[float, float]) -> tuple[float, float]:
        return (pt[0] - self.x0) + self.margin_x, (self.y1 - pt[1]) + self.margin_y


class _LocalTransform(_Transform):
    """`_Transform` anchored at 0,0 instead of a whole slide's own margin:
    no minimum-size clamp (that's a presentation-wide requirement) and no
    margin. `_PlacedTransform`, below, is the version actually used by
    `insert_into_pptx()`, anchored at the target region's own origin
    instead of 0,0."""

    def __init__(self, bbox: tuple[float, float, float, float]):
        self.x0, self.y0, self.x1, self.y1 = bbox
        self.margin_x = self.margin_y = 0.0


class _PlacedTransform(_LocalTransform):
    """`_LocalTransform`, anchored at `origin` (left, top, in slide
    inches) instead of 0,0 -- `insert_into_pptx()`'s own transform, so
    each shape lands directly at the position a wrapping group would
    have put it, with no group actually involved (docs/spec.md SS4.2,
    ext: pikslide never creates a PowerPoint group anywhere -- see
    `_add_all_shapes` -- since PowerPoint's own group-resize math was
    found to silently distort a group's *children*'s sizes)."""

    def __init__(self, bbox: tuple[float, float, float, float], origin: tuple[float, float]):
        super().__init__(bbox)
        self.margin_x, self.margin_y = origin


def _set_arrowheads(line, larrow: bool, rarrow: bool) -> None:
    """python-pptx has no high-level arrowhead API; add the OOXML elements
    directly. <a:headEnd> is the line's start, <a:tailEnd> its end."""
    ln = line._get_or_add_ln()
    for tag, present in (("a:headEnd", larrow), ("a:tailEnd", rarrow)):
        el = ln.find(qn(tag))
        if el is None:
            el = ln.makeelement(qn(tag), {})
            ln.append(el)
        el.set("type", "triangle" if present else "none")


def _apply_line_style(line, shape: Shape) -> None:
    if shape.sw < 0:
        line.fill.background()
        return
    _apply_color(line.color, shape.color or Color(rgb=0))
    line.width = Pt(max(shape.sw, 0.001) * 72)
    if shape.dashed > 0:
        line.dash_style = MSO_LINE_DASH_STYLE.DASH
    elif shape.dotted > 0:
        line.dash_style = MSO_LINE_DASH_STYLE.ROUND_DOT


def _apply_run_font(run, flags: list[str], typeface: str) -> None:
    """Set a run's font family (docs/spec.md SS3.3): by default the
    theme's own minor font (`+mn-lt`/`+mn-ea`), or its major (heading)
    font (`+mj-lt`/`+mj-ea`) when the `major` text flag is present --
    kept as a symbolic theme reference, exactly like a `theme` color
    (never resolved to a literal family here, checked: round-trips and
    renders correctly through real PowerPoint, picking each script's own
    family, not just the Latin one), unless `typeface` (empty by default,
    §3.3/§3.7) is a literal family, which then overrides both the Latin
    and East Asian slots.

    python-pptx's `Font.name` only ever touches `<a:latin>` -- `<a:ea>` has
    no public API at all (checked), so it's set directly on the run's `rPr`."""
    major = "major" in flags
    if typeface:
        latin_val = ea_val = typeface
    else:
        latin_val = "+mj-lt" if major else "+mn-lt"
        ea_val = "+mj-ea" if major else "+mn-ea"
    run.font.name = latin_val
    rPr = run.font._rPr
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rPr.find(qn("a:latin")).addnext(ea)  # <a:ea> must follow <a:latin> in schema order
    ea.set("typeface", ea_val)


def _apply_text(pptx_shape, shape: Shape) -> None:
    """Draws `shape.texts` at `shape.text_sizes`/`.typeface` -- captured
    when this object was created (docs/spec.md SS3.3, ext; `pik/layout.py`
    `_layout_object()`), not read from some document-wide default here,
    so a later `medium`/`typeface` override elsewhere can't retroactively
    change already-drawn text."""
    if not shape.texts:
        return
    tf = pptx_shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    # PowerPoint's default autoshape text margins (~0.1in sides, ~0.05in
    # top/bottom) eat into the box width _autosize_text() already sized to
    # the text, which can force unwanted wrapping. pikchr's own charWidth
    # padding is the only margin that's meant to apply, and that's already
    # folded into _autosize_text()'s width/height estimate.
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, (text, flags) in enumerate(shape.texts):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = (
            PP_ALIGN.LEFT if "ljust" in flags else PP_ALIGN.RIGHT if "rjust" in flags else PP_ALIGN.CENTER
        )
        run = para.add_run()
        run.text = text
        _apply_run_font(run, flags, shape.typeface)
        run.font.bold = "bold" in flags
        run.font.italic = "italic" in flags
        run.font.size = _font_size(flags, shape.text_sizes)
        _apply_color(run.font.color, shape.color or Color(rgb=0))


#: The OOXML SVG-picture extension (Microsoft's own, not ECMA-376): a
#: {GUID} extension on <a:blip> naming a *second* blip, the real SVG,
#: which a modern viewer shows in place of the <a:blip> it's attached to
#: (the PNG fallback older ones use) -- checked against real PowerPoint,
#: rendering the SVG (a crisp vector circle), not the fallback (a solid
#: rectangle) purposely made to look different, so this is really being
#: read, not just tolerated as unrecognized markup.
_SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def _attach_svg_extension(picture, svg_path: str) -> None:
    """Add `svg_path`'s own bytes to the package as an image part, and
    reference it from `picture`'s <a:blip> as the SVG extension (docs/
    spec.md SS3.5) -- `picture` itself was already added from a *rasterized
    PNG* (see _add_image_shape()), since python-pptx cannot add an SVG at
    all (add_picture() raises TypeError, checked: it reads the image with
    Pillow, which can't read SVG either). The SVG part is added the same
    way, bypassing python-pptx's Image/ImagePart helper (also Pillow-
    based) with a raw ImagePart instead -- content type "image/svg+xml"
    needs no image dimensions read from it at all, only its bytes."""
    with open(svg_path, "rb") as f:
        svg_bytes = f.read()
    package = picture.part.package
    svg_part = ImagePart(package.next_image_partname("svg"), "image/svg+xml", package, svg_bytes)
    svg_rId = picture.part.relate_to(svg_part, RT.IMAGE)
    ext_xml = (
        f"<a:extLst {nsdecls('a', 'r')}>"
        f'<a:ext uri="{_SVG_EXT_URI}">'
        f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" r:embed="{svg_rId}"/>'
        "</a:ext>"
        "</a:extLst>"
    )
    picture._element.blipFill.blip.append(parse_xml(ext_xml))


def _add_image_shape(container, shape: Shape, tf: _Transform, name_prefix: str = "") -> None:
    """`image` (docs/spec.md SS3.5). A Picture has no text_frame of its own
    in python-pptx (checked, like a connector), so any text on it is drawn
    as a separate textbox, centred over it -- one box, not per-string
    floating labels the way a line's text is (SS3.5: "a separate text box
    centred on it", singular).

    An SVG is embedded together with a PNG fallback for older viewers
    (see _attach_svg_extension()): add_picture() itself gets the
    rasterized PNG (rasterize_svg(), the same rsvg-convert path
    layout.py's own sizing already uses), then the real SVG is attached
    to it afterward."""
    assert shape.image_path is not None
    left, top, w, h = tf.rect(shape)
    w, h = max(w, 0.01), max(h, 0.01)
    if shape.image_path.lower().endswith(".svg"):
        png_bytes = rasterize_svg(shape.image_path)
        picture = container.shapes.add_picture(
            io.BytesIO(png_bytes), Inches(left), Inches(top), width=Inches(w), height=Inches(h)
        )
        _attach_svg_extension(picture, shape.image_path)
    else:
        picture = container.shapes.add_picture(shape.image_path, Inches(left), Inches(top), width=Inches(w), height=Inches(h))
    if shape.name:
        picture.name = f"{name_prefix}{shape.name}"
    if shape.alt_text:
        # python-pptx 1.0.2 has no real `alt_text` property (checked: it
        # silently becomes a plain, never-saved instance attribute --
        # `pic.alt_text = "..."` raises nothing and even reads back
        # correctly in memory, but the saved file's `descr` is untouched).
        # The actual OOXML attribute is `p:cNvPr/@descr`; add_picture()
        # already sets it to the filename, so this overrides that default.
        picture._element.nvPicPr.cNvPr.set("descr", shape.alt_text)

    if not shape.texts:
        return
    textbox = container.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
    text_frame = textbox.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    for i, (text, flags) in enumerate(shape.texts):
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = text
        _apply_run_font(run, flags, shape.typeface)
        run.font.bold = "bold" in flags
        run.font.italic = "italic" in flags
        run.font.size = _font_size(flags, shape.text_sizes)
        _apply_color(run.font.color, shape.color or Color(rgb=0))


def _add_block_shape(container, shape: Shape, tf: _Transform, name_prefix: str = "") -> None:
    left, top, w, h = tf.rect(shape)
    w, h = max(w, 0.01), max(h, 0.01)
    autoshape_type = _AUTOSHAPE.get(shape.kind, MSO_SHAPE.RECTANGLE)
    if shape.kind == "shape":
        # Validated already, in layout.py's _resolve_preset_name(); the
        # roundRect preset is *literally* MSO_SHAPE.ROUNDED_RECTANGLE
        # (checked), so the rad -> adjustments[0] branch just below applies
        # to `shape roundRect` exactly as it does to `box rad ...`.
        assert shape.preset is not None
        autoshape_type = MSO_SHAPE.from_xml(shape.preset)
    elif shape.kind == "box" and shape.rad > 0:
        autoshape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    pptx_shape = container.shapes.add_shape(autoshape_type, Inches(left), Inches(top), Inches(w), Inches(h))
    if shape.name:
        pptx_shape.name = f"{name_prefix}{shape.name}"
    if autoshape_type == MSO_SHAPE.ROUNDED_RECTANGLE and shape.rad > 0:
        # `adjustments[0]` is the corner radius as a fraction of min(w, h),
        # not an absolute length. `rad == 0` (no explicit `rad` attribute,
        # for either `box rad ...` or a bare `shape roundRect`) leaves
        # python-pptx's own sensible default (checked: ~1/6) alone, rather
        # than flattening a plain `shape roundRect`'s corners to square.
        pptx_shape.adjustments[0] = min(0.5, shape.rad / min(w, h))

    if shape.kind == "text":
        pptx_shape.fill.background()
        pptx_shape.line.fill.background()
    else:
        if shape.fill is None:
            pptx_shape.fill.background()
        else:
            pptx_shape.fill.solid()
            _apply_color(pptx_shape.fill.fore_color, shape.fill)
        _apply_line_style(pptx_shape.line, shape)

    _apply_text(pptx_shape, shape)


def _add_line_shape(container, shape: Shape, tf: _Transform, name_prefix: str = "") -> None:
    assert shape.path is not None
    points = [tf.point(p) for p in shape.path]

    if len(points) == 2:
        (x1, y1), (x2, y2) = points
        connector = container.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        if shape.name:
            connector.name = f"{name_prefix}{shape.name}"
        _apply_line_style(connector.line, shape)
        _set_arrowheads(connector.line, shape.larrow, shape.rarrow)
    else:
        x0, y0 = points[0]
        builder = container.shapes.build_freeform(Inches(x0), Inches(y0))
        builder.add_line_segments([(Inches(x), Inches(y)) for x, y in points[1:]], close=shape.closed)
        freeform = builder.convert_to_shape()
        if shape.name:
            freeform.name = f"{name_prefix}{shape.name}"
        freeform.fill.background()
        _apply_line_style(freeform.line, shape)
        _set_arrowheads(freeform.line, shape.larrow, shape.rarrow)

    _add_line_text(container, shape, tf, name_prefix)


def _line_label_rects(shape: Shape) -> list[tuple[tuple[float, float, float, float], str, list[str]]]:
    """Compute each text label's (x0, y0, x1, y1) rect in pik space (y-up,
    unmargined), alongside its text/flags -- shared by _add_line_text()
    (which draws these) and _content_bbox() (which needs to know how far
    they extend beyond shape.bbox, since a line's own bbox -- just its
    path -- doesn't account for labels floating above/below it). Sized at
    `shape.text_sizes` -- this line's own, as it stood when written
    (docs/spec.md SS3.3, ext) -- not some other, possibly later, size."""
    if not shape.texts:
        return []
    base_size_pt = shape.text_sizes.get("medium", 10.5 / 72.0) * 72.0
    label_box_h = base_size_pt * 1.15 / 72.0  # a touch taller than line_height(), just for rendering safety
    label_step = base_size_pt / 9.0 * _LABEL_STEP_IN
    bx0, by0, bx1, by1 = shape.bbox
    cx = (bx0 + bx1) / 2
    cy = (by0 + by1) / 2
    out = []
    for (text, flags), slot in zip(shape.texts, assign_text_slots(shape.texts)):
        dy = _SLOT_STEP.get(slot, 0) * label_step
        box_w = max(len(text) * base_size_pt / 72.0 * 0.7, 0.3)
        y = cy + dy
        out.append(((cx - box_w / 2, y - label_box_h / 2, cx + box_w / 2, y + label_box_h / 2), text, flags))
    return out


def _add_line_text(container, shape: Shape, tf: _Transform, name_prefix: str = "") -> None:
    """A connector/freeform shape has no text_frame in python-pptx, so a
    line's text (e.g. an arrow's label) is rendered as small floating
    textboxes instead, placed above/on/below the line per
    assign_text_slots(); named "<line name> text <k>" (docs/spec.md SS3.1,
    ext), 1-based among *this line's own* labels, not a diagram-wide count."""
    for i, ((x0, y0, x1, y1), text, flags) in enumerate(_line_label_rects(shape), start=1):
        left, top = tf.point((x0, y1))
        textbox = container.shapes.add_textbox(Inches(left), Inches(top), Inches(x1 - x0), Inches(y1 - y0))
        if shape.name:
            textbox.name = f"{name_prefix}{shape.name} text {i}"
        text_frame = textbox.text_frame
        text_frame.word_wrap = False
        text_frame.margin_left = text_frame.margin_right = 0
        text_frame.margin_top = text_frame.margin_bottom = 0
        para = text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = text
        _apply_run_font(run, flags, shape.typeface)
        run.font.bold = "bold" in flags
        run.font.italic = "italic" in flags
        run.font.size = _font_size(flags, shape.text_sizes)
        _apply_color(run.font.color, shape.color or Color(rgb=0))


def _content_bbox(result: LayoutResult) -> tuple[float, float, float, float]:
    """result.bbox, expanded to also cover line labels -- a line's own
    bbox is just its path, so a label floating above/below it (see
    _line_label_rects()) can stick out past result.bbox on its own.
    flatten_shapes(), not result.shapes directly, since a line can be
    nested inside a block (docs/spec.md SS3.1, ext) at any depth."""
    x0, y0, x1, y1 = result.bbox
    for shape in flatten_shapes(result.shapes):
        if shape.kind not in ("line", "arrow", "spline", "arc"):
            continue
        for (lx0, ly0, lx1, ly1), _text, _flags in _line_label_rects(shape):
            x0, y0 = min(x0, lx0), min(y0, ly0)
            x1, y1 = max(x1, lx1), max(y1, ly1)
    return x0, y0, x1, y1


# ---------------------------------------------------------------------------
# Template settings files (docs/spec.md SS3.8)
# ---------------------------------------------------------------------------


def find_settings_file(target_path: str, explicit: str | None = None) -> str | None:
    """A template's or deck's settings file (docs/spec.md SS3.8):
    `explicit` (`--settings FILE`) if given -- an error if it doesn't
    exist, and nothing is then looked for beside `target_path` -- else
    `<name>.theme.pik` beside `target_path`, or `None` if there isn't
    one. `target_path` is whatever `--template` names, or the OUTPUT deck
    a diagram is being inserted into; the path rules of SS3.6 (`include`'s
    containment) do not apply to either case,
    since this is the tool finding its own file, or the caller naming one
    directly, never a `.pik` naming one itself."""
    if explicit is not None:
        if not os.path.isfile(explicit):
            raise LayoutError(f"--settings file not found: {explicit}")
        return explicit
    stem, dot, _ext = target_path.rpartition(".")
    candidate = f"{stem}.theme.pik" if dot else f"{target_path}.theme.pik"
    return candidate if os.path.isfile(candidate) else None


def write_pptx(
    result: LayoutResult,
    path: str,
    margin: float = 0.15,
) -> None:
    """Render `result` (from :func:`pikslide.pik.layout.resolve_layout`, or
    `resolve_for_pptx`) to a single-slide PowerPoint file at `path`, sized
    to fit the diagram. Each shape draws with its *own* captured text
    size/family (docs/spec.md SS3.3, ext -- `Shape.text_sizes`/`.typeface`,
    as they stood when that object was written), not a single
    document-wide default; empty `typeface`, by far the common case, means
    a symbolic theme font reference rather than a literal name (see
    `_apply_run_font`), so a fresh standalone deck's own built-in Office
    theme decides the actual family, the same as any other new PowerPoint
    file.

    `margin` only needs to cover the diagram's own edge (e.g. a thick
    stroke's outer half, or PowerPoint's arrowhead overshoot) -- line
    labels are already accounted for by _content_bbox(), not by margin."""
    tf = _Transform(_content_bbox(result), margin)
    prs = Presentation()
    prs.slide_width = Emu(int(tf.slide_width * EMU_PER_INCH))
    prs.slide_height = Emu(int(tf.slide_height * EMU_PER_INCH))
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    _add_all_shapes(slide, result.shapes, tf)
    prs.save(path)


# ---------------------------------------------------------------------------
# --template (docs/spec.md SS3.3 rule 2, SS4.1): a new standalone deck
# starting from another file's theme, instead of the built-in Office one.
# ---------------------------------------------------------------------------


def _normalize_potx(path: str) -> str:
    """python-pptx refuses to open a `.potx` as-is (`ValueError`, checked)
    -- the only actual difference from a `.pptx` is the content type
    declared for `/ppt/presentation.xml` (checked: `...template.main+xml`
    vs. `...presentation.main+xml`, nothing else). Rewrites just that and
    saves the result to a new temp file, whose path the caller is
    responsible for removing (`_open_template_base` does)."""
    with zipfile.ZipFile(path) as zin:
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
    new_content_types = content_types.replace(
        "presentationml.template.main+xml", "presentationml.presentation.main+xml"
    )
    if new_content_types == content_types:
        raise LayoutError(f"{path}: doesn't look like a .potx (no template content type found)")
    fd, tmp_path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = new_content_types.encode("utf-8")
            zout.writestr(item, data)
    return tmp_path


def _open_template_base(path: str) -> tuple[Presentation, str | None]:
    """Open `path` (`.pptx` or `.potx`) as a base for a new standalone
    deck: every existing slide removed (docs/spec.md SS3.3: "stripped of
    its sample slides" -- generalized here to any `--template`, not just
    a bare `.potx`, since `--template` may just as well name a real,
    populated deck, and standalone output is one new slide, not the
    template's own N plus one). Returns the temp file path to remove
    afterward too, for a normalized `.potx` (`None` for a plain `.pptx`,
    opened directly, nothing to clean up)."""
    if not os.path.isfile(path):
        raise LayoutError(f"--template file not found: {path}")
    tmp_path = _normalize_potx(path) if path.lower().endswith(".potx") else None
    prs = Presentation(tmp_path if tmp_path is not None else path)
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.get(qn("r:id")))
        sldIdLst.remove(sldId)
    return prs, tmp_path


def _resolve_template_layout(prs: Presentation, layout_name: str):
    """The slide layout a new slide is made from (docs/spec.md SS3.8
    `layout`), which also fixes the master and so the theme (SS3.3):
    named, found in any master (the first match wins if more than one
    master has a layout of that name); empty, the first `blank`-type
    layout of the *first* master, or that master's own first layout if it
    has none (checked: a real-world template's layouts commonly don't set
    the `type` attribute at all, so this fallback -- not just the
    blank-type lookup -- is the common case in practice, not a rare
    corner)."""
    all_layouts = [(m, layout) for m in prs.slide_masters for layout in m.slide_layouts]
    if not all_layouts:
        raise LayoutError("this template has no slide masters/layouts at all")
    if not layout_name:
        first_master = prs.slide_masters[0]
        for layout in first_master.slide_layouts:
            if layout.element.get("type") == "blank":
                return layout
        return first_master.slide_layouts[0]
    for _master, layout in all_layouts:
        if layout.name == layout_name:
            return layout
    known = sorted({layout.name for _master, layout in all_layouts})
    raise LayoutError(f"no slide layout named {layout_name!r} in this template; it has: {', '.join(known)}")


def write_pptx_from_template(
    result: LayoutResult,
    template_path: str,
    path: str,
    layout_name: str = "",
    margin: float = 0.15,
) -> None:
    """`write_pptx()`, but starting from `template_path`'s theme (docs/
    spec.md SS3.3 rule 2, SS4.1 `--template`) instead of the built-in
    Office theme, so theme colors and fonts resolve the way the real
    deck will -- both are emitted symbolically either way (never resolved
    to a literal here, see `_apply_run_font`/`Color`), so this needs no
    theme file of its own read at all: the diagram's new slide is simply
    made from the *right* slide layout (`layout_name`, SS3.8) of
    `template_path` itself, so its symbols resolve against that
    template's own theme once the file is reopened, the same way
    `insert_into_pptx()`'s do against the deck it inserts into.

    The output slide is still sized to the diagram, not the template's
    own slide size (SS4.1): only the theme (via the resolved layout's
    master) comes along, plus the layout's own placeholder shapes, if it
    has any and `layout_name` named it explicitly (the default, empty
    `layout_name` resolves to a *blank* layout precisely to avoid that)."""
    prs, tmp_path = _open_template_base(template_path)
    try:
        layout = _resolve_template_layout(prs, layout_name)
        tf = _Transform(_content_bbox(result), margin)
        prs.slide_width = Emu(int(tf.slide_width * EMU_PER_INCH))
        prs.slide_height = Emu(int(tf.slide_height * EMU_PER_INCH))
        slide = prs.slides.add_slide(layout)
        _add_all_shapes(slide, result.shapes, tf)
        prs.save(path)
    finally:
        if tmp_path is not None:
            os.remove(tmp_path)


def _add_all_shapes(container, shapes: list[Shape], tf: _Transform, name_prefix: str = "") -> None:
    """Draw every shape in `shapes` into `container` (a Slide, for
    `write_pptx()`, `write_pptx_from_template()`, or `insert_into_pptx()`
    alike -- all expose the same `.shapes.add_X()` API, checked, so this
    needs no branching on which one it got).

    pikslide never creates a PowerPoint group anywhere in its own output
    (ext): PowerPoint's own group-resize math was found to silently
    distort a group's *children*'s sizes. A "block" shape (docs/spec.md
    SS3.1: "Blocks are groups" at the *pik* language level) is therefore
    flattened here -- recursed straight into the same `container`/`tf`,
    not nested inside a group of its own -- so its children land as
    ordinary, individually-selectable top-level shapes, each still
    keeping its own name (`_layout_statements()` has already given every
    shape, block or not, one: its label, or a default "<class> <n>"); the
    block shape itself has no PowerPoint shape representing it at all. No
    local/nested transform is needed either: a block's children already
    sit in the *same* global pik coordinate space as everything else
    (`_translate()`, at layout time, already placed them there), so `tf`
    -- the one already in effect for this whole call -- maps them
    correctly without adjustment.

    `name_prefix`, non-empty only for `insert_into_pptx()` (docs/spec.md
    SS4.2, ext), is prepended to every shape's own name there, so a later
    re-run can find and remove exactly (and only) what a previous run
    under the same prefix placed, leaving anything else on the slide --
    including an edit a person made by hand since -- untouched."""
    for shape in shapes:
        if shape.kind == "block":
            _add_all_shapes(container, shape.sublist, tf, name_prefix)
        elif shape.kind in NOT_RENDERED:
            continue
        elif shape.kind in ("line", "arrow", "spline", "arc"):
            _add_line_shape(container, shape, tf, name_prefix)
        elif shape.kind == "image":
            _add_image_shape(container, shape, tf, name_prefix)
        else:
            _add_block_shape(container, shape, tf, name_prefix)


# ---------------------------------------------------------------------------
# Inserting into an existing deck (docs/spec.md SS4.2)
# ---------------------------------------------------------------------------


def _find_shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def _is_empty_placeholder(shape) -> bool:
    return shape.is_placeholder and (not shape.has_text_frame or not shape.text_frame.text.strip())


def _shape_rect_inches(shape) -> tuple[float, float, float, float]:
    return shape.left.inches, shape.top.inches, shape.width.inches, shape.height.inches


def _shapes_union_rect(shapes: list) -> tuple[float, float, float, float]:
    """(left, top, width, height) in inches of the smallest rectangle
    covering every shape in `shapes` (non-empty) -- used in place of a
    single group's own rect now that a diagram's shapes are never
    wrapped in one (see `_add_all_shapes`)."""
    lefts, tops, rights, bottoms = zip(
        *(
            (sh.left.inches, sh.top.inches, sh.left.inches + sh.width.inches, sh.top.inches + sh.height.inches)
            for sh in shapes
        )
    )
    left, top = min(lefts), min(tops)
    return left, top, max(rights) - left, max(bottoms) - top


def _resolve_region(
    slide,
    region: str | None,
    rect: tuple[float, float, float, float] | None,
    previous: list,
    content_area: tuple[float, float, float, float] | None,
):
    """Return ((left, top, width, height) in inches, the shape to delete
    afterward if it was an empty placeholder or else None) for
    `--region`/`--rect` (docs/spec.md SS4.2). With neither, `content_area`
    (a settings file's `content_left`/etc., SS3.8) is the default target;
    with none of the three, an error. `previous` is whatever shapes a
    prior run of `insert_into_pptx()` under the same prefix placed on
    this slide (possibly none), used as a fallback below."""
    if region is not None and rect is not None:
        raise LayoutError("--region and --rect are mutually exclusive")
    if region is not None:
        shape = _find_shape_by_name(slide, region)
        if shape is None:
            # Idempotent re-runs (SS4.2): the *first* run already deleted
            # an empty placeholder used as the region, so it can't be
            # found by name a second time. Fall back to the extent of
            # whatever this tool placed there before, if anything -- its
            # own union rect is exactly "where the diagram lives now".
            if previous:
                return _shapes_union_rect(previous), None
            names = [sh.name for sh in slide.shapes if sh.name]
            raise LayoutError(f"no shape named {region!r} on this slide{_did_you_mean(region, names)}")
        to_delete = shape if _is_empty_placeholder(shape) else None
        return _shape_rect_inches(shape), to_delete
    if rect is not None:
        return rect, None
    if content_area is not None:
        return content_area, None
    raise LayoutError(
        "no target region: pass --region or --rect, or use a template's "
        "settings file to give it a content area (docs/spec.md SS3.8)"
    )


#: --align choices (ext -- docs/spec.md SS4.2 says only that a smaller
#: diagram defaults to "the region's top-left" and that --align
#: "overrides", without enumerating values; the standard 3x3 grid needs no
#: further design decision beyond that default). Each value is (fx, fy),
#: the fraction of the region's own extra room (region size - diagram
#: size) to place before the diagram on each axis.
ALIGN_CHOICES = {
    "top-left": (0.0, 0.0), "top-center": (0.5, 0.0), "top-right": (1.0, 0.0),
    "center-left": (0.0, 0.5), "center": (0.5, 0.5), "center-right": (1.0, 0.5),
    "bottom-left": (0.0, 1.0), "bottom-center": (0.5, 1.0), "bottom-right": (1.0, 1.0),
}


def _align_offset(align: str, region_w: float, region_h: float, diagram_w: float, diagram_h: float) -> tuple[float, float]:
    if align not in ALIGN_CHOICES:
        raise LayoutError(f"unknown --align {align!r}; choices: {', '.join(sorted(ALIGN_CHOICES))}")
    fx, fy = ALIGN_CHOICES[align]
    return (region_w - diagram_w) * fx, (region_h - diagram_h) * fy


def insert_into_pptx(
    result: LayoutResult,
    deck_path: str,
    slide_no: int,
    region: str | None = None,
    rect: tuple[float, float, float, float] | None = None,
    group_id: str = "diagram",
    align: str = "top-left",
    prefix: str | None = None,
) -> Presentation:
    """Insert `result` into slide `slide_no` (1-based) of the deck at
    `deck_path` (docs/spec.md SS4.2). Returns the modified `Presentation`
    -- the caller saves it, to a copy or in place; unlike `write_pptx()`,
    this doesn't save directly, since there's an existing file whose path
    (a copy, or the same one) is the caller's call, not this function's.

    Every shape is added directly to the slide, named `<prefix><shape's
    own name>` -- never wrapped in a PowerPoint group (ext: see
    `_add_all_shapes`). `prefix` defaults to `pik:<group_id>`; pass it
    explicitly (`--prefix`) to override that outright, when `group_id`
    alone (e.g. a short, ordinary-looking id like "arch") would risk
    colliding with a real shape name already on the slide.

    A diagram is never scaled (SS4.2): one larger than its region is an
    error, not silently shrunk. A smaller one sits at the region's
    top-left by default; `align` (one of `ALIGN_CHOICES`) overrides.
    Running this again against the same `deck_path`/`slide_no`/`prefix`
    finds every shape whose name starts with `prefix`, removes exactly
    those (keeping their former z-order position for the replacements),
    and adds the new ones -- so a second run replaces what pikslide
    itself placed, rather than adding a second copy, while anything else
    on the slide (including an edit a person made by hand since the last
    run) is left completely untouched.

    With neither `region` nor `rect`, `result.content_area` (a settings
    file's content area, SS3.8) is the default target.

    Text is drawn with each shape's own `.typeface`/`.text_sizes`
    (docs/spec.md SS3.3, ext): empty typeface, the common case, means a
    symbolic theme font reference (`+mn-lt` etc., see `_apply_run_font`)
    rather than a literal name -- which needs no theme file of its own to
    be read here at all, since the reference resolves against whatever
    theme `deck_path`'s own slide master already uses once the file is
    back open in PowerPoint (SS3.3 rule 1)."""
    prs = Presentation(deck_path)
    if not 1 <= slide_no <= len(prs.slides):
        raise LayoutError(f"--slide {slide_no} is out of range: this deck has {len(prs.slides)} slide(s)")
    slide = prs.slides[slide_no - 1]
    name_prefix = prefix if prefix is not None else f"pik:{group_id}"
    previous = [sh for sh in slide.shapes if sh.name and sh.name.startswith(name_prefix)]

    (region_left, region_top, region_w, region_h), to_delete = _resolve_region(
        slide, region, rect, previous, result.content_area
    )

    bbox = _content_bbox(result)
    diagram_w, diagram_h = bbox[2] - bbox[0], bbox[3] - bbox[1]
    if diagram_w > region_w + 1e-6 or diagram_h > region_h + 1e-6:
        raise LayoutError(
            f"diagram ({diagram_w:.2f}in x {diagram_h:.2f}in) is larger than its "
            f"region ({region_w:.2f}in x {region_h:.2f}in); a diagram is never "
            "scaled (docs/spec.md SS4.2) -- adjust the source instead"
        )
    dx, dy = _align_offset(align, region_w, region_h, diagram_w, diagram_h)

    spTree = slide.shapes._spTree
    insert_index = None
    if previous:
        insert_index = min(list(spTree).index(sh._element) for sh in previous)
        for sh in previous:
            spTree.remove(sh._element)

    before = set(spTree)
    tf = _PlacedTransform(bbox, (region_left + dx, region_top + dy))
    _add_all_shapes(slide, result.shapes, tf, name_prefix)
    added = [el for el in spTree if el not in before]

    if insert_index is not None:
        for el in added:
            spTree.remove(el)
        for offset, el in enumerate(added):
            spTree.insert(insert_index + offset, el)

    if to_delete is not None:
        spTree.remove(to_delete._element)

    return prs
