"""pikslide as a library (docs/spec.md SS6, issue #21): parse(), layout()
and render_into() onto a slide the caller has."""

from __future__ import annotations

import pathlib

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches

import pikslide


def _slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[5])  # Title Only


def _template_deck(tmp_path: pathlib.Path) -> pathlib.Path:
    path = tmp_path / "tmpl.pptx"
    Presentation().save(str(path))
    return path


def test_render_into_draws_on_the_slide_beside_what_it_has():
    _prs, slide = _slide()
    result = pikslide.layout(pikslide.parse('Web: box "Web"\narrow "go"\nDb: box "Db"\n'))
    names = pikslide.render_into(slide, result, origin=(Inches(1), Inches(2)))
    assert names == {"Web": "Web", "arrow 1": "arrow 1", "Db": "Db"}
    assert [s.name for s in slide.shapes] == ["Title 1", "Web", "arrow 1", "arrow 1 text 1", "Db"]


def test_render_into_puts_the_top_left_of_what_is_drawn_at_origin():
    # The arrow's label is above the boxes: it is the top.
    _prs, slide = _slide()
    text = 'box "A" ht 0.2in\narrow 0.8in "a much longer label than the arrow" "on two lines" above\nbox "B" ht 0.2in\n'
    result = pikslide.layout(pikslide.parse(text))
    pikslide.render_into(slide, result, origin=(Inches(1), Inches(2)))
    drawn = list(slide.shapes)[1:]
    left = min(Emu(s.left).inches for s in drawn)
    top = min(Emu(s.top).inches for s in drawn)
    right = max(Emu(s.left + s.width).inches for s in drawn)
    bottom = max(Emu(s.top + s.height).inches for s in drawn)
    assert (left, top) == pytest.approx((1.0, 2.0), abs=0.001)
    # Its size is the box round the object list's: the boxes and the
    # arrow's label, whose points are inside them.
    boxes = [b for o in result.objects() for b in ([o] if "x" in o else o["labels"])]
    xs = [x for b in boxes for x in (b["x"] - b["wid"] / 2, b["x"] + b["wid"] / 2)]
    ys = [y for b in boxes for y in (b["y"] - b["ht"] / 2, b["y"] + b["ht"] / 2)]
    assert (right - left, bottom - top) == pytest.approx((max(xs) - min(xs), max(ys) - min(ys)), abs=0.01)


def test_a_name_the_slide_has_moves_on_past_the_diagrams_own():
    # Spec SS3.1: the diagram's Start-2 keeps its name, so its Start,
    # which the slide has, moves on to Start-3.
    _prs, slide = _slide()
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1)).name = "Start"
    result = pikslide.layout(pikslide.parse('[ Start: box "a" ]\n[ Start: box "b" ]\n'))
    assert [o["name"] for o in result.objects()] == ["Start", "Start-2"]
    names = pikslide.render_into(slide, result, origin=(0, 0))
    assert names == {"Start": "Start-3", "Start-2": "Start-2"}


def test_a_lines_label_text_boxes_count_as_its_names():
    _prs, slide = _slide()
    slide.shapes.add_textbox(0, 0, Inches(1), Inches(1)).name = "Go text 1"
    names = pikslide.render_into(slide, pikslide.layout(pikslide.parse('Go: arrow "go"\n')), origin=(0, 0))
    assert names == {"Go": "Go-2"}
    assert [s.name for s in slide.shapes][-2:] == ["Go-2", "Go-2 text 1"]


def test_name_suffix_goes_after_every_name_and_a_taken_one_still_moves_on():
    _prs, slide = _slide()
    result = pikslide.layout(pikslide.parse('Start: box "a"\narrow\n'))
    assert pikslide.render_into(slide, result, origin=(0, 0), name_suffix="flow") == {
        "Start": "Start-flow", "arrow 1": "arrow 1-flow"}
    assert pikslide.render_into(slide, result, origin=(0, 0), name_suffix="flow") == {
        "Start": "Start-flow-2", "arrow 1": "arrow 1-flow-2"}


def test_render_into_attaches_lines_on_the_callers_slide():
    _prs, slide = _slide()
    result = pikslide.layout(pikslide.parse('A: box\narrow attach\nB: box\n'))
    pikslide.render_into(slide, result, origin=(0, 0))
    a, arrow, b = list(slide.shapes)[1:]
    assert arrow.begin_x == a.left + a.width and arrow.end_x == b.left


def test_origin_is_in_emu():
    _prs, slide = _slide()
    result = pikslide.layout(pikslide.parse('box\n'))
    with pytest.raises(TypeError, match="EMU"):
        pikslide.render_into(slide, result, origin=(1.0, 2.0))


def test_layout_warns_as_check_does_by_the_names_in_the_object_list():
    long = '"a string much too wide for the box it is in"'
    result = pikslide.layout(pikslide.parse(f"[ A: box {long} ]\n[ A: box {long} ]\n"))
    overflows = [w for w in result.warnings() if "overflows" in w["message"]]
    assert [w["objects"] for w in overflows] == [["A"], ["A-2"]]
    assert result.warnings()[-1]["message"].startswith("no template given")


def test_layout_with_a_template_leaves_out_the_warnings_about_its_slide(tmp_path: pathlib.Path):
    # A diagram wider than the Title and Content layout's content area
    # warns with --template, but render_into()'s caller places it.
    tmpl = _template_deck(tmp_path)
    result = pikslide.layout(pikslide.parse('box "A" width 9.5\n'), template=str(tmpl))
    assert result.warnings() == []


def test_layout_reads_the_templates_settings_file(tmp_path: pathlib.Path):
    tmpl = _template_deck(tmp_path)
    (tmp_path / "tmpl.theme.pik").write_text("boxwid = 2in\n", encoding="utf-8")
    [box] = pikslide.layout(pikslide.parse('box "A"\n'), template=str(tmpl)).objects()
    assert box["wid"] == 2.0


def test_slide_layout_picks_the_layout_and_so_the_master(tmp_path: pathlib.Path):
    # The default layout never raises this: the name reaches the template.
    tmpl = _template_deck(tmp_path)
    doc = pikslide.parse('box "A"\n')
    assert pikslide.layout(doc, template=str(tmpl), slide_layout="Title Only").objects()
    with pytest.raises(pikslide.LayoutError, match="no slide layout named 'Nope'"):
        pikslide.layout(doc, template=str(tmpl), slide_layout="Nope")
    with pytest.raises(ValueError, match="needs template"):
        pikslide.layout(doc, slide_layout="Title Only")


def test_an_image_resolves_against_the_base_dir_parse_was_given(tmp_path: pathlib.Path):
    Image.new("RGB", (100, 50)).save(tmp_path / "logo.png")
    _prs, slide = _slide()
    result = pikslide.layout(pikslide.parse('image "logo.png" width 1\n', base_dir=str(tmp_path)))
    pikslide.render_into(slide, result, origin=(0, 0))
    assert [s.name for s in slide.shapes][-1] == "image 1"


def test_errors_are_the_packages_own():
    with pytest.raises(pikslide.PikSyntaxError):
        pikslide.parse('"unterminated\n')
    with pytest.raises(pikslide.LayoutError):
        pikslide.layout(pikslide.parse("box fill nosuchcolor\n"))
