"""Tests for attached lines (issue #18): `attach` and `path`, from parsing
to the connector's attached ends in the deck and --check's object list."""

from __future__ import annotations

import json
import pathlib

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from pikslide import main
from pikslide.pik import ast, parse
from pikslide.pik.layout import LayoutError, connection_sites, flatten_shapes, resolve_layout
from pikslide.pik.tokens import PikSyntaxError
from pikslide.pptx_writer import resolve_for_pptx, write_pptx


def layout(src: str):
    return resolve_layout(parse(src))


def attached(src: str, name: str | None = None) -> list[tuple[str, int] | None]:
    """The objects and sites the last line `src` draws (or `name`) attached its ends to."""
    lines = [s for s in flatten_shapes(layout(src).shapes) if s.attach]
    line = next(s for s in lines if s.name == name) if name else lines[-1]
    return [None if end is None else (end[0].name, end[1]) for end in line.attached]


def warnings(src: str) -> list[str]:
    return [w.message for w in layout(src).warnings]


def test_attach_and_path_parse():
    [stmt] = parse("arrow from A.e to B.w attach path straight\n").statements
    assert stmt.attributes[2:] == [ast.Attach(), ast.ConnectorPath("straight")]
    with pytest.raises(PikSyntaxError, match="after 'path'"):
        parse("arrow attach path\n")


def test_connection_sites_are_the_presets_in_a_connectors_order():
    # A rectangle's: top, left, bottom, right. An ellipse's eight go round
    # counterclockwise from the top, on its compass points.
    box = layout("box\n").shapes[0]
    assert connection_sites(box) == [pytest.approx(p) for p in [(0, 0.25), (-0.375, 0), (0, -0.25), (0.375, 0)]]
    circle = layout("circle\n").shapes[0]
    compass = ["n", "nw", "w", "sw", "s", "se", "e", "ne"]
    assert connection_sites(circle) == [pytest.approx(circle.edge_point(e)) for e in compass]
    assert connection_sites(layout("Q: (0, 0)\n").shapes[0]) == []


def test_ends_attach_to_the_objects_from_and_to_name():
    src = "A: box\nB: circle at (2, 1)\n"
    assert attached(src + "arrow from A.e to B.w attach\n") == [("A", 3), ("B", 2)]
    assert attached(src + "arrow from A.n to B.nw attach\n") == [("A", 0), ("B", 1)]
    assert warnings(src + "arrow from A.e to B.w attach\n") == []


def test_an_unnamed_end_attaches_to_the_object_whose_site_it_is_on_placed_before_or_after():
    assert attached("A: box\narrow attach\nB: box\n") == [("A", 3), ("B", 1)]


def test_an_end_off_a_site_warns_with_the_nearest_and_stays_unattached():
    src = "A: box\nB: box at (2, 0)\n"
    assert attached(src + "arrow from A.ne to B.w attach\n") == [None, ("B", 1)]
    assert warnings(src + "arrow from A.ne to B.w attach\n") == [
        "arrow 1's start isn't on a connection site of A, so it isn't attached: end it at A.e, the nearest"]
    assert attached(src + "arrow from A.e + (0, 0.1) to B.w attach\n") == [None, ("B", 1)]
    # A default name can't be written in the source as it is.
    assert "end it at the .w of box 2" in warnings("box\nbox at (2, 0)\narrow from 1st box.e to 2nd box.nw attach\n")[0]


@pytest.mark.parametrize("src, what", [
    ("Q: (0, 0)\nA: box at (2, 0)\narrow from Q to A.w attach\n", "Q, a point"),
    ("M: move\nA: box at (2, 0)\narrow from M.end to A.w attach\n", "M, a move"),
    ("L: line from (0, 0) right 1\nA: box at (2, 0)\narrow from L.end to A.w attach\n", "L, a line"),
])
def test_an_end_on_an_object_with_no_sites_warns(src: str, what: str):
    [message] = warnings(src)
    assert message == f"arrow 1's start is on {what}, which nothing attaches to: it isn't attached"
    assert attached(src) == [None, ("A", 1)]


def test_an_unnamed_end_on_two_objects_sites_warns():
    src = "A: box\nB: box with .w at A.e\nC: box at (3, 0)\narrow from (0.375, 0) to C.w attach\n"
    assert attached(src) == [None, ("C", 1)]
    [message] = warnings(src)
    assert message.startswith("arrow 1's start is on a connection site of both A and B")
    assert [w.objects for w in layout(src).warnings] == [("arrow 1", "A", "B")]


def test_an_unnamed_end_inside_an_object_isnt_on_it():
    # Only its outline counts: an arrow ending inside a frame is fine.
    src = "F: box wid 3 ht 2\nA: box at F\narrow from A.e right 0.5 attach\n"
    assert attached(src) == [("A", 3), None]
    assert warnings(src) == []


def test_a_path_a_straight_connector_cant_take_warns_and_stays_unattached():
    src = "A: box\nB: box at (2, -1)\narrow from A.e right 1 then down until even with B.n attach\n"
    assert attached(src) == [None, None]
    assert warnings(src) == ["arrow 1 has 2 segments, and a straight connector has one: "
                             "it is written as drawn, not attached"]


@pytest.mark.parametrize("src, message", [
    ("box attach", "'attach' is only valid on a line or an arrow, not a box"),
    ("arc attach", "'attach' is only valid on a line or an arrow, not an arc"),
    ("arrow from (0, 0) to (1, 0) attach", "'attach' on arrow 1: neither of its ends is on an object"),
    ("A: box\narrow from A.e path straight", "'path straight' is the form of an attached line's connector: "
                                             "it needs 'attach'"),
    ("A: box\narrow from A.e attach path wavy", "unknown connector form 'path wavy': use straight, elbow or curved"),
])
def test_attach_errors(src: str, message: str):
    with pytest.raises(LayoutError, match="^" + message.replace("(", r"\(").replace(")", r"\)") + "$"):
        layout(src + "\n")


def render(src: str, tmp_path: pathlib.Path) -> dict[str, object]:
    out = tmp_path / "out.pptx"
    write_pptx(resolve_for_pptx(parse(src), base_dir=str(tmp_path)), str(out))
    return {s.name: s for s in Presentation(str(out)).slides[0].shapes}


def ends(connector, shapes) -> list[tuple[str, int] | None]:
    names = {s.shape_id: name for name, s in shapes.items()}
    nv = connector._element.nvCxnSpPr.cNvCxnSpPr
    return [None if (el := nv.find(qn(tag))) is None else (names[int(el.get("id"))], int(el.get("idx")))
            for tag in ("a:stCxn", "a:endCxn")]


def test_an_attached_line_is_a_connector_attached_where_it_was_drawn(tmp_path: pathlib.Path):
    # B is drawn after the arrow, and the circle's site 1 is its .nw.
    shapes = render("A: box\narrow attach\nB: box\nC: circle at (2, 1)\narrow from A.n to C.nw attach\n", tmp_path)
    assert ends(shapes["arrow 1"], shapes) == [("A", 3), ("B", 1)]
    assert ends(shapes["arrow 2"], shapes) == [("A", 0), ("C", 1)]
    # Not moved: still where pikslide drew it.
    plain = render("A: box\narrow\nB: box\nC: circle at (2, 1)\narrow from A.n to C.nw\n", tmp_path)
    for name in ("arrow 1", "arrow 2"):
        a, b = shapes[name], plain[name]
        assert (a.begin_x, a.begin_y, a.end_x, a.end_y) == (b.begin_x, b.begin_y, b.end_x, b.end_y)
        assert ends(b, plain) == [None, None]


def test_a_picture_and_a_text_take_an_attached_end(tmp_path: pathlib.Path):
    from PIL import Image

    Image.new("RGB", (40, 20), "white").save(tmp_path / "x.png")
    shapes = render('I: image "x.png" wid 0.8\nT: text "note" at (0, -1)\narrow from T.n to I.s attach\n', tmp_path)
    assert ends(shapes["arrow 1"], shapes) == [("T", 0), ("I", 2)]


def test_check_says_which_ends_are_attached(monkeypatch, capsys, tmp_path: pathlib.Path):
    src = tmp_path / "d.pik"
    src.write_text("A: box\nB: box at (2, 0)\narrow from A.e to B.w attach\narrow from A.ne to B.w attach\n",
                   encoding="utf-8")
    monkeypatch.setattr("sys.argv", ["pikslide", str(src), "--check"])
    main()
    objects = {o["name"]: o for o in json.loads(capsys.readouterr().out)["objects"]}
    assert objects["arrow 1"]["attached"] == {"start": "A", "end": "B"}
    assert objects["arrow 2"]["attached"] == {"end": "B"}
    assert "attached" not in objects["A"]
