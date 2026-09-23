"""`pikslide.main()` -- the `argparse` CLI (docs/spec.md SS4)."""

from __future__ import annotations

import json
import pathlib

import pytest
from pptx import Presentation

from pikslide import main

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _write(tmp_path: pathlib.Path, name: str, text: str) -> pathlib.Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def _run(monkeypatch, argv: list[str]) -> None:
    monkeypatch.setattr("sys.argv", ["pikslide", *argv])
    main()


# ---------------------------------------------------------------------------
# Basic rendering
# ---------------------------------------------------------------------------


def test_no_args_prints_hello(monkeypatch, capsys):
    _run(monkeypatch, [])
    assert "Hello from pikslide" in capsys.readouterr().out


def test_dump_mode_prints_parsed_tree(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    _run(monkeypatch, [str(src)])
    out = capsys.readouterr().out
    assert "Document" in out and "box" in out.lower()


def test_standalone_render_to_pptx(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out_path = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out_path)])
    assert out_path.exists()
    assert f"wrote {out_path}" in capsys.readouterr().out


def test_unsupported_output_format_errors(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    with pytest.raises(SystemExit) as exc:
        _run(monkeypatch, [str(src), str(tmp_path / "d.svg")])
    assert exc.value.code == 1
    assert "unsupported output format" in capsys.readouterr().err


def test_syntax_error_prints_message_and_exits(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", '"unterminated\n')
    with pytest.raises(SystemExit) as exc:
        _run(monkeypatch, [str(src)])
    assert exc.value.code == 1
    assert "error:" in capsys.readouterr().err


# ---------------------------------------------------------------------------
# --template / --settings (docs/spec.md SS3.3 rule 1/2, SS3.8, SS4)
# ---------------------------------------------------------------------------


def _template_deck(tmp_path: pathlib.Path, name: str = "tmpl.pptx") -> pathlib.Path:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Sample"
    path = tmp_path / name
    prs.save(str(path))
    return path


def test_standalone_without_template_warns_and_still_writes(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out)])
    captured = capsys.readouterr()
    assert out.exists()
    assert "warning:" in captured.err and "stand-ins" in captured.err
    assert "wrote" in captured.out


def test_strict_turns_the_no_template_warning_into_an_error(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    with pytest.raises(SystemExit) as exc:
        _run(monkeypatch, [str(src), str(out), "--strict"])
    assert exc.value.code == 1
    assert not out.exists()
    assert "stand-ins" in capsys.readouterr().err


def test_template_flag_starts_a_new_deck_from_that_theme(monkeypatch, capsys, tmp_path):
    tmpl = _template_deck(tmp_path)
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out), "--template", str(tmpl)])
    captured = capsys.readouterr()
    assert out.exists()
    assert "warning:" not in captured.err  # a real template given -- no stand-in warning
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    assert [s.name for s in prs.slides[0].shapes] == ["box 1"]  # sample slide stripped


def test_output_omitted_defaults_from_input_when_template_is_given(monkeypatch, capsys, tmp_path):
    tmpl = _template_deck(tmp_path)
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    _run(monkeypatch, [str(src), "--template", str(tmpl)])
    default_out = tmp_path / "d.pptx"
    assert default_out.exists()
    assert f"wrote {default_out}" in capsys.readouterr().out


def test_settings_file_beside_template_is_found_automatically(monkeypatch, tmp_path):
    tmpl = _template_deck(tmp_path)
    _write(tmp_path, "tmpl.theme.pik", 'typeface = "Verdana"\n')
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out), "--template", str(tmpl)])
    run = Presentation(str(out)).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Verdana"


def test_explicit_settings_flag_overrides_the_beside_file(monkeypatch, tmp_path):
    tmpl = _template_deck(tmp_path)
    _write(tmp_path, "tmpl.theme.pik", 'typeface = "Verdana"\n')
    _write(tmp_path, "elsewhere.pik", 'typeface = "Georgia"\n')
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out), "--template", str(tmpl), "--settings", str(tmp_path / "elsewhere.pik")])
    run = Presentation(str(out)).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Georgia"


def test_missing_explicit_settings_file_is_an_error(monkeypatch, tmp_path):
    tmpl = _template_deck(tmp_path)
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    with pytest.raises(SystemExit) as exc:
        _run(monkeypatch, [str(src), str(tmp_path / "out.pptx"), "--template", str(tmpl), "--settings", str(tmp_path / "nope.pik")])
    assert exc.value.code == 1


# ---------------------------------------------------------------------------
# --include-path, --check, --format json (ext)
# ---------------------------------------------------------------------------


def test_include_path_flag_is_searched_when_not_found_beside_the_source(monkeypatch, tmp_path):
    lib_dir = tmp_path / "lib"
    lib_dir.mkdir()
    (lib_dir / "house.pik").write_text("primary = accent3\n", encoding="utf-8")
    src = _write(tmp_path, "d.pik", 'include "house.pik"\nbox fill primary\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out), "--include-path", str(lib_dir)])
    assert out.exists()


def test_check_parses_and_lays_out_without_writing(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out), "--check"])
    assert not out.exists()
    assert "ok" in capsys.readouterr().out


def test_check_still_catches_a_layout_error(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", "box fill nosuchcolor\n")
    with pytest.raises(SystemExit) as exc:
        _run(monkeypatch, [str(src), str(tmp_path / "d.pptx"), "--check"])
    assert exc.value.code == 1


def test_format_json_success(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", 'box "Web"\n')
    out = tmp_path / "d.pptx"
    _run(monkeypatch, [str(src), str(out), "--format", "json"])
    payload = json.loads(capsys.readouterr().out)
    assert payload["ok"] is True
    assert payload["output"] == str(out)
    assert any("stand-ins" in w for w in payload["warnings"])


def test_format_json_syntax_error_has_file_line_column(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", '"unterminated\n')
    with pytest.raises(SystemExit):
        _run(monkeypatch, [str(src), "--format", "json"])
    payload = json.loads(capsys.readouterr().out)
    assert payload["ok"] is False
    err = payload["errors"][0]
    assert err["file"] == str(src)
    assert err["line"] == 1
    assert isinstance(err["column"], int)


def test_format_json_layout_error_has_no_position(monkeypatch, capsys, tmp_path):
    src = _write(tmp_path, "d.pik", "box fill nosuchcolor\n")
    out = tmp_path / "d.pptx"
    with pytest.raises(SystemExit):
        _run(monkeypatch, [str(src), str(out), "--format", "json"])
    payload = json.loads(capsys.readouterr().out)
    err = payload["errors"][0]
    assert err["line"] is None
    assert "nosuchcolor" in err["message"]
