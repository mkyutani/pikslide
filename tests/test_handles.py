"""Tests for a preset shape's handles (issue #20): `adjust NAME VALUE` and
`tip at POSITION`, from parsing to the adjust values in the deck."""

from __future__ import annotations

import pathlib
import re

import pytest
from pptx import Presentation

from pikslide.pik import ast, parse
from pikslide.pik.layout import LayoutError, flatten_shapes, resolve_layout
from pikslide.pik.text_area import text_rect
from pikslide.pptx_writer import resolve_for_pptx, write_pptx


def shape(src: str, name: str | None = None):
    """The last shape `src` draws, or the one called `name`."""
    shapes = flatten_shapes(resolve_layout(parse(src)).shapes)
    return next(s for s in shapes if s.name == name) if name else shapes[-1]


def adjust(src: str, name: str | None = None) -> dict[str, int]:
    return {k: round(v) for k, v in shape(src, name).adjust.items()}


def test_adjust_and_tip_parse():
    [stmt] = parse("shape rightArrow adjust head.length 0.3 adjust shaft.width 40% adjust adj1 -8333 "
                   "tip at (1, 2)\n").statements
    assert stmt.attributes == [
        ast.Adjust("head.length", ast.RelExpr(abs=ast.Num(0.3))),
        ast.Adjust("shaft.width", ast.RelExpr(percent=ast.Num(40.0))),
        ast.Adjust("adj1", ast.RelExpr(abs=ast.UnaryOp("-", ast.Num(8333.0)))),
        ast.Tip(ast.Coord(ast.Num(1.0), ast.Num(2.0))),
    ]


def test_a_named_handle_is_a_length_or_a_percentage_of_its_side():
    # 0.75 x 0.5: the head is measured in the shorter side, the shaft in the height.
    assert adjust("shape rightArrow adjust head.length 0.15 adjust shaft.width 40%\n") == {
        "adj2": 30000, "adj1": 40000}
    assert adjust("shape rightArrow adjust head.length 20% adjust shaft.width 0.1\n") == {
        "adj2": 30000, "adj1": 20000}
    # An upright arrow's shaft is measured in its width, its head along its height.
    assert adjust("shape upArrow wid 0.5 ht 1 adjust shaft.width 0.25 adjust head.length 50%\n") == {
        "adj1": 50000, "adj2": 100000}
    assert adjust("shape chevron adjust head.length 0.1\n") == {"adj": 20000}


def test_a_handle_by_the_presets_own_name_takes_the_value_as_written():
    assert adjust("shape bentArrow adjust adj1 30000 adjust adj4 -5\n") == {"adj1": 30000, "adj4": -5}
    assert adjust("shape can adjust adj 10000\n") == {"adj": 10000}


@pytest.mark.parametrize(("src", "message"), [
    ("box adjust adj 1\n", "'adjust' is only valid on a preset shape"),
    ("shape rect adjust adj 1\n", "rect has no handles to adjust"),
    ("shape rightArrow adjust head.lenght 0.2\n",
     "rightArrow has no handle head.lenght: it has shaft.width, head.length, adj1, adj2; did you mean head.length"),
    ("shape rightArrow adjust adj1 50%\n", "adjust adj1 takes the value OOXML writes, 100000 for 1, not a %"),
    ("shape rightArrow adjust head.length -0.1\n", "adjust head.length can't be negative"),
    ("box tip at (1, 1)\n", "'tip' is only valid on a preset shape"),
    ("shape rightArrow tip at (1, 1)\n", "'tip' is only valid on a callout"),
])
def test_handle_errors(src: str, message: str):
    with pytest.raises(LayoutError, match=re.escape(message)):
        resolve_layout(parse(src))


def test_a_handle_longer_than_the_preset_draws_warns():
    result = resolve_layout(parse("A: shape leftRightArrow adjust head.length 0.5\n"))
    [w] = result.warnings
    assert w.message.startswith("A: head.length 0.50 in is longer than leftRightArrow can draw it (0.38 in")
    assert w.objects == ("A",)
    assert resolve_layout(parse("shape leftRightArrow adjust head.length 0.375\n")).warnings == []


def test_tip_points_a_wedge_callout():
    # From the center, in the shape's width and height, y down.
    assert adjust("shape wedgeRectCallout at (0, 0) tip at (1.5, -1)\n") == {"adj1": 200000, "adj2": 200000}
    assert adjust("B: box at (2, 1)\nshape cloudCallout at (0, 0) tip at B.w\n") == {"adj1": 216667, "adj2": -200000}


def test_tip_ends_a_leader_line_leaving_the_side_facing_it():
    # To the right, level with the box's middle: from the right side, straight.
    assert adjust("shape callout1 at (0, 0) tip at (1, 0)\n") == {
        "adj1": 50000, "adj2": 108333, "adj3": 50000, "adj4": 183333}
    # Above and to the left, but further above: from the top, at its left corner.
    assert adjust("shape borderCallout1 at (0, 0) tip at (-0.5, 1)\n") == {
        "adj1": -8333, "adj2": 0, "adj3": -150000, "adj4": -16667}
    # A bend out from the side; two bends, then along it to the tip's level.
    assert adjust("shape callout2 at (0, 0) tip at (-1.5, -0.5)\n") == {
        "adj1": 100000, "adj2": -8333, "adj3": 100000, "adj4": -16667, "adj5": 150000, "adj6": -150000}
    assert adjust("shape callout3 at (0, 0) tip at (-1.5, -0.5)\n") == {
        "adj1": 100000, "adj2": -8333, "adj3": 100000, "adj4": -16667, "adj5": 150000, "adj6": -16667,
        "adj7": 150000, "adj8": -150000}
    # Straight below: from the bottom, at the tip's x.
    assert adjust("shape callout3 at (0, 0) tip at (0.1, -1)\n") == {
        "adj1": 108333, "adj2": 63333, "adj3": 116667, "adj4": 63333, "adj5": 116667, "adj6": 63333,
        "adj7": 250000, "adj8": 63333}


def test_an_accent_callouts_line_leaves_a_left_or_right_side():
    # Its accent bar is upright at the line's start, so not from the bottom.
    values = adjust("shape accentCallout1 at (0, 0) tip at (0.1, -1)\n")
    assert values["adj2"] == 108333


def test_handles_stay_relative_to_a_shape_a_block_moves():
    alone = adjust("A: box\nshape wedgeRectCallout at A + (1, 1) tip at A.ne\n")
    moved = adjust("[ A: box\nshape wedgeRectCallout at A + (1, 1) tip at A.ne ] at (5, 5)\n", "shape 1")
    assert moved == alone


def test_same_as_copies_the_adjust_values_of_the_same_preset():
    src = "A: shape rightArrow adjust head.length 0.1\n"
    assert adjust(src + "shape rightArrow same as A\n") == {"adj2": 20000}
    assert adjust(src + "shape leftArrow same as A\n") == {}


def test_fit_and_the_overflow_warning_follow_the_handles():
    src = 'shape rightArrow "0123456789" fit adjust head.length 0.4 adjust shaft.width 60%\n'
    s, box = shape(src), shape('shape rect "0123456789" fit\n')
    left, top, right, bottom = text_rect("rightArrow", s.w, s.h, s.adjust)
    assert s.adjust["adj2"] == pytest.approx(0.4 / min(s.w, s.h) * 100000)
    assert right - left >= box.w - 1e-6 and bottom - top >= box.h - 1e-6
    # A longer head leaves less room for the text: 1 in of it fits in 2 in
    # less a 0.25 in head, but not less 1.8 in.

    class FixedMetrics:
        def text_width(self, text, flags=(), text_sizes=None, typeface=""):
            return 0.1 * len(text)

        def line_height(self, flags=(), text_sizes=None):
            return 0.2

    def warnings(src: str) -> list[str]:
        return [w.message for w in resolve_layout(parse(src), metrics=FixedMetrics()).warnings]

    assert warnings('shape rightArrow "0123456789" wid 2\n') == []
    [w] = warnings('shape rightArrow "0123456789" wid 2 adjust head.length 1.8\n')
    assert "text overflows the usable width (1.00 in > 0.90 in)" in w


def test_the_deck_has_the_adjust_values(tmp_path: pathlib.Path):
    out = tmp_path / "out.pptx"
    write_pptx(resolve_for_pptx(parse(
        "shape rightArrow adjust head.length 0.15\n"
        'shape roundRect "x" rad 0.1 width 1 height 1 adjust adj 30000\n'
        "shape wedgeRectCallout at (0, 0) tip at (1.5, -1)\n"
    )), str(out))
    arrow, card, callout = Presentation(str(out)).slides[0].shapes

    def guides(s):
        return [(gd.get("name"), gd.get("fmla")) for gd in s._element.spPr.prstGeom.iter()
                if gd.tag.endswith("}gd")]

    assert guides(arrow) == [("adj2", "val 30000")]
    assert guides(card) == [("adj", "val 30000")]
    assert guides(callout) == [("adj1", "val 200000"), ("adj2", "val 200000")]
