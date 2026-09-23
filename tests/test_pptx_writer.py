"""Tests for pikslide.pptx_writer: rendering a resolved layout to PowerPoint."""

from __future__ import annotations

import pathlib
import zipfile

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL_TYPE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from pikslide.pik import parse
from pikslide.pik.layout import LayoutError
from pikslide.pptx_writer import resolve_for_pptx, write_pptx

FIXTURES_DIR = pathlib.Path(__file__).parent / "fixtures" / "examples"
EXAMPLE_FILES = sorted(FIXTURES_DIR.glob("*.pik"))


def render(text: str, tmp_path: pathlib.Path, name: str = "out.pptx") -> Presentation:
    # Uses resolve_for_pptx() (real font-metrics-based "fit" sizing), the
    # same path the CLI takes, so these tests exercise it too. base_dir
    # defaults to tmp_path itself, so an `image "x.png"` resolves against
    # whatever the test already wrote there.
    result = resolve_for_pptx(parse(text), base_dir=str(tmp_path))
    out = tmp_path / name
    write_pptx(result, str(out))
    return Presentation(str(out))


@pytest.mark.parametrize("path", EXAMPLE_FILES, ids=lambda p: p.name)
def test_official_example_renders_without_error(path: pathlib.Path, tmp_path: pathlib.Path):
    prs = render(path.read_text(encoding="utf-8"), tmp_path, path.stem + ".pptx")
    assert len(prs.slides) == 1


def test_box_becomes_a_shape_with_text(tmp_path: pathlib.Path):
    prs = render('box "Hello"\n', tmp_path)
    slide = prs.slides[0]
    assert len(slide.shapes) == 1
    shape = slide.shapes[0]
    assert shape.text_frame.text == "Hello"


def test_two_point_arrow_becomes_a_connector(tmp_path: pathlib.Path):
    prs = render("arrow right\n", tmp_path)
    slide = prs.slides[0]
    assert len(slide.shapes) == 1
    assert slide.shapes[0].shape_type == MSO_SHAPE_TYPE.LINE


def test_arrow_labels_render_as_textboxes_above_and_below(tmp_path: pathlib.Path):
    # A connector shape has no text_frame of its own in python-pptx, so an
    # arrow's text must show up as separate floating textboxes -- and two
    # un-flagged texts split above/below the line (assign_text_slots()).
    prs = render('arrow right "Top" "Bottom"\n', tmp_path)
    slide = prs.slides[0]
    textboxes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert {tb.text_frame.text for tb in textboxes} == {"Top", "Bottom"}
    top_box = next(tb for tb in textboxes if tb.text_frame.text == "Top")
    bottom_box = next(tb for tb in textboxes if tb.text_frame.text == "Bottom")
    # slide y grows downward, so the "above the line" label has the smaller top.
    assert top_box.top < bottom_box.top


def test_rounded_box_gets_rounded_rectangle_with_scaled_corner(tmp_path: pathlib.Path):
    prs = render("box rad 0.1 width 1 height 1\n", tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.adjustments[0] == pytest.approx(0.1, abs=0.01)


def test_multipoint_line_becomes_a_freeform(tmp_path: pathlib.Path):
    prs = render("line right 1 then up 1\n", tmp_path)
    slide = prs.slides[0]
    assert len(slide.shapes) == 1
    assert slide.shapes[0].shape_type == MSO_SHAPE_TYPE.FREEFORM


def test_slide_size_matches_diagram_bbox_plus_margin(tmp_path: pathlib.Path):
    result = resolve_for_pptx(parse("box\n"))
    out = tmp_path / "sized.pptx"
    write_pptx(result, str(out), margin=0.5)
    prs = Presentation(str(out))
    x0, y0, x1, y1 = result.bbox
    assert prs.slide_width.inches == pytest.approx((x1 - x0) + 1.0, abs=0.01)
    assert prs.slide_height.inches == pytest.approx((y1 - y0) + 1.0, abs=0.01)


def test_fill_color_is_applied(tmp_path: pathlib.Path):
    # A color name is an ordinary lowercase variable, defined by the
    # prelude (docs/spec.md SS2, SS3.7) -- not a capitalized PLACENAME.
    prs = render("box fill red\n", tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.fill.fore_color.rgb == RGBColor(0xFF, 0x00, 0x00)


def test_invis_object_has_no_outline(tmp_path: pathlib.Path):
    prs = render("box invis\n", tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.line.fill.type == MSO_FILL_TYPE.BACKGROUND


def test_theme_color_is_emitted_as_schemeclr_not_resolved_rgb(tmp_path: pathlib.Path):
    # docs/spec.md SS3.3: emitted as schemeClr, never as resolved RGB, so
    # swapping the theme restyles the diagram.
    prs = render('box fill accent1 lighter 40%\n', tmp_path)
    shape = prs.slides[0].shapes[0]
    fore = shape.fill.fore_color
    assert fore.theme_color == MSO_THEME_COLOR.ACCENT_1
    assert fore.brightness == pytest.approx(0.4)


def test_no_fill_leaves_the_shape_transparent(tmp_path: pathlib.Path):
    # The prelude's own default (`fill = none`, docs/spec.md SS3.7).
    prs = render('box "hi"\n', tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.fill.type == MSO_FILL_TYPE.BACKGROUND


def test_text_size_flag_selects_the_large_point_size(tmp_path: pathlib.Path):
    prs = render('box "small" small\nbox "large" large\n', tmp_path)
    small_run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    large_run = prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    assert small_run.font.size.pt == pytest.approx(9.0)
    assert large_run.font.size.pt == pytest.approx(12.0)


def test_overriding_medium_changes_the_default_text_size(tmp_path: pathlib.Path):
    prs = render('medium = 14pt\nbox "hi"\n', tmp_path)
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(14.0)


def test_a_mid_document_override_does_not_retroactively_resize_earlier_text(tmp_path: pathlib.Path):
    # docs/spec.md SS3.3, ext -- checked as the actual bug this was
    # before Shape.text_sizes existed: both ended up at 20pt.
    prs = render('box "a" small\nsmall = 20pt\nbox "b" small\n', tmp_path)
    a_run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    b_run = prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    assert a_run.font.size.pt == pytest.approx(9.0)
    assert b_run.font.size.pt == pytest.approx(20.0)


def test_a_mid_document_typeface_override_does_not_retroactively_change_earlier_text(tmp_path: pathlib.Path):
    prs = render('box "a"\ntypeface = "Georgia"\nbox "b"\n', tmp_path)
    a_run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    b_run = prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    assert a_run.font.name == "+mn-lt"
    assert b_run.font.name == "Georgia"


def test_fit_sizing_tracks_a_medium_override_at_render_time_too(tmp_path: pathlib.Path):
    # Same fix, checked through the real PilFontMetrics-backed render()
    # path (tests/test_layout.py's equivalent test uses a stub metrics
    # instead, to isolate the mechanism from font-file availability).
    small = render('box "Wide Wide Wide Text" fit\n', tmp_path, name="small.pptx").slides[0].shapes[0]
    big = render('medium = 30pt\nbox "Wide Wide Wide Text" fit\n', tmp_path, name="big.pptx").slides[0].shapes[0]
    assert big.width.inches > small.width.inches
    assert big.height.inches > small.height.inches


# ---------------------------------------------------------------------------
# Fonts (docs/spec.md SS3.3): symbolic theme references by default, `major`,
# `typeface` -- checked, as elsewhere, against what's actually saved to disk
# (not just what's readable back in memory before saving).
# ---------------------------------------------------------------------------


def _ea_typeface(run) -> str | None:
    ea = run.font._rPr.find(qn("a:ea"))
    return ea.get("typeface") if ea is not None else None


def test_default_font_is_the_symbolic_minor_theme_reference(tmp_path: pathlib.Path):
    prs = render('box "hi"\n', tmp_path)
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "+mn-lt"
    assert _ea_typeface(run) == "+mn-ea"


def test_major_flag_selects_the_major_theme_reference(tmp_path: pathlib.Path):
    prs = render('box "Heading" major\n', tmp_path)
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "+mj-lt"
    assert _ea_typeface(run) == "+mj-ea"


def test_typeface_variable_overrides_with_a_literal_family(tmp_path: pathlib.Path):
    prs = render('typeface = "Comic Sans MS"\nbox "hi" major\n', tmp_path)
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    # Overrides both slots, and applies regardless of `major` (docs/spec.md
    # SS3.3: "a specific family" -- one value for everything, not a
    # major/minor pair of its own).
    assert run.font.name == "Comic Sans MS"
    assert _ea_typeface(run) == "Comic Sans MS"


def test_typeface_must_be_a_string(tmp_path: pathlib.Path):
    with pytest.raises(LayoutError):
        resolve_for_pptx(parse("typeface = 5\nbox\n"))


def test_line_label_and_image_caption_also_get_the_theme_font(tmp_path: pathlib.Path):
    # _add_line_text() and _add_image_shape() build runs independently of
    # _apply_text() -- checked separately, since nothing shares that code path.
    _make_image(tmp_path)
    prs = render('arrow right "Label"\nimage "logo.png" width 50% "Caption"\n', tmp_path, name="fonts.pptx")
    slide = prs.slides[0]
    label = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and s.text_frame.text == "Label")
    caption = next(s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and s.text_frame.text == "Caption")
    for box in (label, caption):
        run = box.text_frame.paragraphs[0].runs[0]
        assert run.font.name == "+mn-lt"
        assert _ea_typeface(run) == "+mn-ea"


# ---------------------------------------------------------------------------
# Object identity: names, groups, z-order (docs/spec.md SS3.1, ext)
# ---------------------------------------------------------------------------


def test_no_shape_inherits_the_theme_effect_style(tmp_path: pathlib.Path):
    # python-pptx's <p:style> points each shape at a theme effect style
    # (effectRef idx 2 for a shape, 1 for a connector), which in many real
    # templates is a drop shadow; an empty <a:effectLst/> in spPr, placed
    # after <a:ln> per the schema, overrides it.
    slide = render('box "A"\narrow\nline right then down\ntext "t"\n', tmp_path).slides[0]
    styled = [s for s in slide.shapes if s._element.find(qn("p:style")) is not None]
    assert len(styled) == 4
    for s in styled:
        sp_pr = s._element.spPr
        tags = [child.tag for child in sp_pr]
        assert qn("a:effectLst") in tags, s.name
        assert len(sp_pr.find(qn("a:effectLst"))) == 0
        assert tags.index(qn("a:effectLst")) > tags.index(qn("a:ln"))


def test_shape_names_are_set_on_the_saved_pptx_shapes(tmp_path: pathlib.Path):
    prs = render('box "a"\nWeb: box "b"\n', tmp_path)
    names = [s.name for s in prs.slides[0].shapes]
    assert names == ["box 1", "Web"]


def test_line_and_image_shapes_are_named_too(tmp_path: pathlib.Path):
    _make_image(tmp_path)
    prs = render('arrow right\nimage "logo.png" width 1\n', tmp_path, name="named.pptx")
    names = [s.name for s in prs.slides[0].shapes]
    assert names == ["arrow 1", "image 1"]


def test_line_label_textbox_naming(tmp_path: pathlib.Path):
    prs = render('Conn: arrow right "Top" "Bottom"\n', tmp_path)
    boxes = {s.text_frame.text: s.name for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX}
    assert boxes["Top"] in ("Conn text 1", "Conn text 2")
    assert boxes["Bottom"] in ("Conn text 1", "Conn text 2")
    assert boxes["Top"] != boxes["Bottom"]


def test_block_is_flattened_not_grouped(tmp_path: pathlib.Path):
    # A "block" ([ ... ], docs/spec.md SS3.1) is never rendered as a
    # PowerPoint group (ext: PowerPoint's own group-resize math was found
    # to silently distort a group's children's sizes) -- its children
    # become ordinary, individually-selectable top-level shapes instead;
    # the block's own label ("Outer") has no PowerPoint shape of its own.
    prs = render('Outer: [ A: box "x"; B: box "y" ]\n', tmp_path)
    top = prs.slides[0].shapes
    assert [s.name for s in top] == ["A", "B"]
    assert all(s.shape_type != MSO_SHAPE_TYPE.GROUP for s in top)


def test_nested_blocks_are_flattened_with_correct_positions(tmp_path: pathlib.Path):
    prs = render(
        'box "before"\n'
        'Outer: [\n'
        '  A: box "a"\n'
        '  Inner: [ D: box "d" ]\n'
        ']\n',
        tmp_path,
    )
    top = prs.slides[0].shapes
    # No group anywhere, at any nesting depth -- every shape, however
    # deeply nested in the source, ends up a flat, top-level sibling.
    assert [s.name for s in top] == ["box 1", "A", "D"]
    assert all(s.shape_type != MSO_SHAPE_TYPE.GROUP for s in top)
    before, a, d = top
    # "a"/"d" sit where the layout placed them in the shared, global pik
    # coordinate space -- to the right of "before", not at some
    # block-local origin (there is no local frame to translate out of
    # anymore, since there's no group to have positioned them relative to).
    assert a.left.inches > before.left.inches + before.width.inches - 0.01
    assert d.left.inches > before.left.inches + before.width.inches - 0.01


def test_behind_places_the_shape_earlier_in_z_order(tmp_path: pathlib.Path):
    # Z-order is draw order (docs/spec.md SS3.1): the shape added *first*
    # ends up visually behind one added later, so "C behind A" must put C
    # before A in the saved shape list.
    prs = render('A: box "a"\nB: box "b"\nC: box "c" behind A\n', tmp_path)
    assert [s.name for s in prs.slides[0].shapes] == ["C", "A", "B"]


def test_preset_shape_renders_as_the_named_autoshape(tmp_path: pathlib.Path):
    prs = render('shape chevron "Step 1"\n', tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.auto_shape_type == MSO_SHAPE.CHEVRON
    assert shape.text_frame.text == "Step 1"


def test_bare_round_rect_preset_keeps_pptx_own_default_corner(tmp_path: pathlib.Path):
    # Checked: python-pptx's own fresh-shape default is ~1/6, not 0 -- a
    # bare `shape roundRect` (no explicit `rad`) must not flatten that.
    prs = render('shape roundRect "x"\n', tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    assert shape.adjustments[0] == pytest.approx(0.16667, abs=0.001)


def test_explicit_rad_on_a_round_rect_preset_overrides_the_corner(tmp_path: pathlib.Path):
    prs = render('shape roundRect "x" rad 0.1 width 1 height 1\n', tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.adjustments[0] == pytest.approx(0.1, abs=0.01)


def _make_image(tmp_path: pathlib.Path, name: str = "logo.png") -> None:
    from PIL import Image

    Image.new("RGB", (400, 200), "white").save(tmp_path / name)


def test_image_renders_as_a_picture(tmp_path: pathlib.Path):
    _make_image(tmp_path)
    prs = render('image "logo.png" width 2\n', tmp_path)
    picture = prs.slides[0].shapes[0]
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert picture.width.inches == pytest.approx(2.0)
    assert picture.height.inches == pytest.approx(1.0)  # 400x200 source, aspect preserved


def test_image_alt_text_is_the_actual_saved_description(tmp_path: pathlib.Path):
    # python-pptx 1.0.2 has no real alt_text property (checked: a plain
    # instance attribute, never serialized) -- this must be the real
    # p:cNvPr/@descr the saved file actually carries.
    _make_image(tmp_path)
    prs = render('image "logo.png" alt "A logo"\n', tmp_path)
    picture = prs.slides[0].shapes[0]
    descr = picture._element.nvPicPr.cNvPr.get("descr")
    assert descr == "A logo"


def test_image_text_becomes_a_centred_caption_textbox(tmp_path: pathlib.Path):
    _make_image(tmp_path)
    prs = render('image "logo.png" width 2 "Caption"\n', tmp_path)
    picture, caption = prs.slides[0].shapes
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert caption.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
    assert caption.text_frame.text == "Caption"
    # centred over the picture, not offset above/below like a line's label.
    assert caption.left == picture.left
    assert caption.top == picture.top


# ---------------------------------------------------------------------------
# SVG images (docs/spec.md SS3.5, ext): a PNG fallback plus the real SVG,
# via the OOXML SVG-picture extension -- checked against real PowerPoint
# separately (it renders the SVG, not the fallback, when they differ).
# ---------------------------------------------------------------------------


def _make_svg(tmp_path: pathlib.Path, name: str = "icon.svg", w: int = 200, h: int = 200) -> pathlib.Path:
    path = tmp_path / name
    path.write_text(
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" '
        f'viewBox="0 0 {w} {h}"><circle cx="{w / 2}" cy="{h / 2}" r="{min(w, h) / 2}" fill="green"/></svg>',
        encoding="utf-8",
    )
    return path


def test_svg_renders_as_a_picture_sized_by_aspect_ratio(tmp_path: pathlib.Path):
    _make_svg(tmp_path, w=400, h=200)
    prs = render('image "icon.svg" width 2\n', tmp_path)
    picture = prs.slides[0].shapes[0]
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE
    assert picture.width.inches == pytest.approx(2.0)
    assert picture.height.inches == pytest.approx(1.0)  # 400x200, aspect preserved


def test_svg_has_a_real_blip_and_the_svg_extension(tmp_path: pathlib.Path):
    _make_svg(tmp_path)
    prs = render('image "icon.svg" width 1\n', tmp_path)
    picture = prs.slides[0].shapes[0]
    blip = picture._element.blipFill.blip
    assert blip.get(qn("r:embed"))  # the PNG fallback -- add_picture()'s normal blip
    ext = blip.find(qn("a:extLst") + "/" + qn("a:ext"))
    assert ext is not None
    assert ext.get("uri") == "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
    svg_blip = ext[0]
    assert svg_blip.tag.endswith("}svgBlip")
    svg_rId = svg_blip.get(qn("r:embed"))
    svg_part = picture.part.related_part(svg_rId)
    assert svg_part.content_type == "image/svg+xml"
    assert svg_part.blob.startswith(b"<svg") or b"<svg" in svg_part.blob[:100]


def test_svg_alt_text_is_the_actual_saved_description(tmp_path: pathlib.Path):
    _make_svg(tmp_path)
    prs = render('image "icon.svg" alt "A circle"\n', tmp_path)
    picture = prs.slides[0].shapes[0]
    assert picture._element.nvPicPr.cNvPr.get("descr") == "A circle"


def test_svg_without_rsvg_convert_is_a_clear_error(tmp_path: pathlib.Path, monkeypatch):
    monkeypatch.setattr("shutil.which", lambda _name: None)
    _make_svg(tmp_path)
    with pytest.raises(LayoutError, match="rsvg-convert"):
        render('image "icon.svg"\n', tmp_path)


# ---------------------------------------------------------------------------
# --template: a new standalone deck from another file's theme
# (docs/spec.md SS3.3 rule 1, SS4, ext)
# ---------------------------------------------------------------------------


def _template_deck(tmp_path: pathlib.Path, name: str = "tmpl.pptx") -> pathlib.Path:
    """A deck with one sample slide (to be stripped) whose default theme
    has the built-in Office layouts, so both the blank-default and a
    named-layout case are reachable."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Sample (should not survive --template)"
    path = tmp_path / name
    prs.save(str(path))
    return path


def _as_potx(pptx_path: pathlib.Path, potx_path: pathlib.Path) -> None:
    """Build a `.potx` by rewriting `pptx_path`'s own content types --
    the reverse of `_normalize_potx()`, so tests don't depend on a real
    `.potx` file existing anywhere on the machine."""
    with zipfile.ZipFile(str(pptx_path)) as zin:
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
    content_types = content_types.replace(
        "presentationml.presentation.main+xml", "presentationml.template.main+xml"
    )
    with zipfile.ZipFile(str(pptx_path)) as zin, zipfile.ZipFile(str(potx_path), "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = content_types.encode("utf-8")
            zout.writestr(item, data)


def _from_template(text: str, template: pathlib.Path, tmp_path: pathlib.Path, name: str = "out.pptx", **kwargs):
    from pikslide.pptx_writer import write_pptx_from_template

    result = resolve_for_pptx(parse(text))
    out = tmp_path / name
    write_pptx_from_template(result, str(template), str(out), **kwargs)
    return Presentation(str(out))


def test_template_strips_the_sample_slide_and_sizes_to_the_diagram(tmp_path: pathlib.Path):
    tmpl = _template_deck(tmp_path)
    prs = _from_template('box "Web"\n', tmpl, tmp_path)
    assert len(prs.slides) == 1
    assert "Sample" not in (prs.slides[0].shapes.title.text if prs.slides[0].shapes.title else "")
    assert prs.slide_width.inches < 2  # sized to the tiny diagram, not the template's own 10x7.5in


def test_template_uses_a_blank_layout_by_default(tmp_path: pathlib.Path):
    tmpl = _template_deck(tmp_path)
    prs = _from_template('box "Web"\n', tmpl, tmp_path)
    names = [s.name for s in prs.slides[0].shapes]
    assert names == ["box 1"]  # no placeholders brought in


def test_template_named_layout_brings_its_placeholders(tmp_path: pathlib.Path):
    tmpl = _template_deck(tmp_path)
    prs = _from_template('box "Web"\n', tmpl, tmp_path, layout_name="Title and Content")
    names = [s.name for s in prs.slides[0].shapes]
    assert "box 1" in names
    assert len(names) > 1  # the layout's own placeholders came along too


def test_template_unknown_layout_name_lists_the_known_ones(tmp_path: pathlib.Path):
    tmpl = _template_deck(tmp_path)
    with pytest.raises(LayoutError, match="Blank"):  # one of the built-in layout names
        _from_template('box "Web"\n', tmpl, tmp_path, layout_name="Nope")


def test_template_colors_and_fonts_stay_symbolic(tmp_path: pathlib.Path):
    tmpl = _template_deck(tmp_path)
    prs = _from_template('box "Web" fill accent1\n', tmpl, tmp_path)
    shape = prs.slides[0].shapes[0]
    assert shape.fill.fore_color.theme_color == MSO_THEME_COLOR.ACCENT_1
    run = shape.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "+mn-lt"


def test_potx_template_is_normalized_and_used(tmp_path: pathlib.Path):
    pptx_path = _template_deck(tmp_path, "tmpl.pptx")
    potx_path = tmp_path / "tmpl.potx"
    _as_potx(pptx_path, potx_path)
    prs = _from_template('box "Web"\n', potx_path, tmp_path)
    assert len(prs.slides) == 1
    assert [s.name for s in prs.slides[0].shapes] == ["box 1"]


def test_template_file_not_found_is_a_clear_error(tmp_path: pathlib.Path):
    with pytest.raises(LayoutError, match="not found"):
        _from_template('box "Web"\n', tmp_path / "nope.pptx", tmp_path)


def test_pptx_misnamed_as_potx_is_a_clear_error(tmp_path: pathlib.Path):
    # A real .pptx (presentation content type already, not template) saved
    # under a .potx name: _normalize_potx()'s replace() is a no-op, so
    # this must be caught rather than silently "succeeding" on an
    # unmodified file that just happens to already open fine.
    pptx_path = _template_deck(tmp_path, "tmpl.pptx")
    fake_potx = tmp_path / "fake.potx"
    fake_potx.write_bytes(pptx_path.read_bytes())
    with pytest.raises(LayoutError, match="potx"):
        _from_template('box "Web"\n', fake_potx, tmp_path)
