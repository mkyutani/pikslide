"""Tests for turning and mirroring a preset shape (issue #26): `rotate DEG`,
`flipH` and `flipV`, from parsing to the deck."""

from __future__ import annotations

import json
import pathlib

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from pikslide import main
from pikslide.pik import ast, parse
from pikslide.pik.layout import LayoutError, callout_points, flatten_shapes, resolve_layout, turn
from pikslide.pptx_writer import resolve_for_pptx, write_pptx


def shapes(src: str):
    return flatten_shapes(resolve_layout(parse(src)).shapes)


def test_rotate_and_flips_parse():
    [stmt] = parse("shape triangle rotate -90 flipH flipV\n").statements
    assert stmt.attributes[1:] == [ast.Flip("h"), ast.Flip("v")]
    [triangle] = shapes("shape triangle rotate -90 flipH flipV\n")
    assert (triangle.rotate, triangle.flip_h, triangle.flip_v) == (270, True, True)


@pytest.mark.parametrize("word", ["rotate 90", "flipH", "flipV"])
def test_only_a_preset_shape_turns(word: str):
    with pytest.raises(LayoutError, match=f"'{word.split()[0]}' is only valid on a preset shape"):
        shapes(f"box {word}\n")


def test_the_box_is_where_the_turned_shape_is_drawn():
    # wid and ht size it unturned; its box, edges and A.wid are as drawn.
    arrow, line = shapes("A: shape rightArrow wid 1 ht 0.5 rotate 90 at (0, 0)\nline from A.s down 0.5\n")
    assert arrow.drawn == (1, 0.5)
    assert (arrow.w, arrow.h) == pytest.approx((0.5, 1))
    assert arrow.bbox == pytest.approx((-0.25, -0.5, 0.25, 0.5))
    assert line.path[0] == pytest.approx((0, -0.5))  # the arrow's point, turned down
    [turned] = shapes("shape rect wid 1 ht 1 rotate 45\n")
    assert (turned.w, turned.h) == pytest.approx((2 ** 0.5, 2 ** 0.5))


def test_turn_flips_then_turns_clockwise():
    # As PowerPoint draws it: rightArrow, flipH then rotate 90, points up.
    [arrow] = shapes("shape rightArrow wid 1 ht 0.5 rotate 90 flipH at (0, 0)\n")
    [head] = turn(arrow, [(0.5, 0)])
    assert head == pytest.approx((0, 0.5))
    assert turn(arrow, [head], back=True) == [pytest.approx((0.5, 0))]


def test_a_turned_callout_still_points_at_its_tip():
    [callout] = shapes("P: (2.5, -1.2)\nshape wedgeRectCallout wid 1 ht 0.6 rotate 30 at (0, 0) tip at P\n")
    assert callout_points(callout) == [pytest.approx((2.5, -1.2))]


def test_a_line_attaches_to_a_turned_shapes_connection_site():
    # A triangle's apex, turned to the right.
    *_, arrow = shapes("T: shape triangle wid 0.6 ht 0.4 rotate 90 at (0, 0)\narrow from (0.2, 0) right 1 attach\n")
    start, _end = arrow.attached
    assert start is not None and start[0].name == "T"


def test_same_as_copies_the_turn():
    *_, copy = shapes("A: shape triangle rotate 45 flipV\nshape triangle same as A\n")
    assert (copy.rotate, copy.flip_h, copy.flip_v) == (45, False, True)


def test_the_deck_has_the_turn_and_flips(tmp_path: pathlib.Path):
    out = tmp_path / "out.pptx"
    write_pptx(resolve_for_pptx(parse("shape triangle rotate -90 flipH\nshape triangle\n")), str(out))
    turned, plain = (s._element.spPr.find(qn("a:xfrm")) for s in Presentation(str(out)).slides[0].shapes)
    assert (turned.get("rot"), turned.get("flipH"), turned.get("flipV")) == (str(270 * 60000), "1", None)
    assert (plain.get("rot"), plain.get("flipH")) == (None, None)
    # Its own size, unturned, which PowerPoint turns about its center.
    ext = turned.find(qn("a:ext"))
    assert (int(ext.get("cx")), int(ext.get("cy"))) == (int(0.75 * 914400), int(0.5 * 914400))


def test_the_slide_holds_a_turned_shape(tmp_path: pathlib.Path):
    out = tmp_path / "out.pptx"
    write_pptx(resolve_for_pptx(parse("shape rect wid 2 ht 0.2 rotate 90\n")), str(out), margin=0)
    prs = Presentation(str(out))
    assert (prs.slide_width.inches, prs.slide_height.inches) == pytest.approx((1, 2))  # 1 in at least


def test_overlap_follows_the_turn():
    def warnings(turned: str) -> list[str]:
        src = f"A: shape rect wid 2 ht 0.2 {turned} at (0, 0)\nB: box wid 0.4 ht 0.4 at (0, 0.7)\n"
        return [w.message for w in resolve_for_pptx(parse(src)).warnings if "overlap" in w.message]

    assert warnings("") == []
    [w] = warnings("rotate 90")
    assert "A" in w and "B" in w


def test_check_reports_the_turn(monkeypatch, capsys, tmp_path: pathlib.Path):
    src = tmp_path / "d.pik"
    src.write_text("M: shape triangle rotate -90 flipV\n", encoding="utf-8")
    monkeypatch.setattr("sys.argv", ["pikslide", str(src), "--check"])
    main()
    [m] = json.loads(capsys.readouterr().out)["objects"]
    assert (m["rotate"], m["flipV"], "flipH" in m) == (270, True, False)
